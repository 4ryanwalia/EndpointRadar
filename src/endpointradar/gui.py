"""Modern dark-themed endpoint security dashboard."""
from __future__ import annotations

import json
import threading
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import messagebox, ttk
from typing import Any

import matplotlib

matplotlib.use("TkAgg")
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg  # noqa: E402
from matplotlib.figure import Figure  # noqa: E402

from . import risk, scanner
from . import threats as threat_analysis
from .reporting import build_threat_word_report, build_word_report

BG = "#121212"
CARD = "#1E1E1E"
TEXT = "#E6E6E6"
MUTED = "#A8A8A8"
ACCENT = "#5DA9E9"
DEVICE_TYPE_COLORS = {
    "Windows": "#64B5F6",
    "Mobile": "#FFB74D",
    "Unknown": "#9E9E9E",
}


class EndpointDashboard(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("Endpoint Security Dashboard")
        self.geometry("1440x920")
        self.minsize(1280, 820)
        self.resizable(True, True)
        try:
            self.state("zoomed")
        except Exception:
            pass
        self.configure(bg=BG)

        self._configure_styles()

        self.results: list[dict[str, Any]] = scanner.load_scan_results()
        self.unauth_results: list[dict[str, Any]] = scanner.load_unauthenticated_scan_results()
        self.devices: list[dict[str, Any]] = []
        self.device_rows: list[dict[str, Any]] = []
        self.selected_device: dict[str, Any] | None = None
        self.selected_scan: dict[str, Any] | None = None
        self.selected_unauth_scan: dict[str, Any] | None = None
        self.connected_devices: dict[str, dict[str, Any]] = {}
        self.live_process_cache: dict[str, list[dict[str, Any]]] = {}
        self.live_process_last_refreshed: dict[str, str] = {}
        self.live_process_fetch_in_progress = False
        self.remote_scan_in_progress = False
        self.unauth_scan_in_progress = False
        self.component_bars: dict[str, ttk.Progressbar] = {}
        self.component_labels: dict[str, tk.Label] = {}
        self.component_score_cache: dict[str, int] = {}
        self.component_contrib_cache: dict[str, float] = {}
        self.subnet_var = tk.StringVar(value=scanner.guess_default_subnet())

        self._build_layout()
        self._refresh_all()

    def _configure_styles(self) -> None:
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure("Dark.TFrame", background=CARD)
        style.configure("Dark.TNotebook", background=CARD, borderwidth=0)
        style.configure("Dark.TNotebook.Tab", background="#2A2A2A", foreground=TEXT, padding=(10, 6))
        style.map("Dark.TNotebook.Tab", background=[("selected", "#333333")], foreground=[("selected", TEXT)])
        style.configure("Risk.Horizontal.TProgressbar", troughcolor="#2B2B2B", background=ACCENT, bordercolor="#2B2B2B")
        style.configure(
            "Dark.Treeview",
            background="#181818",
            foreground=TEXT,
            fieldbackground="#181818",
            borderwidth=0,
            rowheight=26,
            font=("Segoe UI", 10),
        )
        style.configure(
            "Dark.Treeview.Heading",
            background=CARD,
            foreground=TEXT,
            relief=tk.FLAT,
            font=("Segoe UI", 10, "bold"),
        )

    def _card(self, parent: tk.Widget) -> tk.Frame:
        frame = tk.Frame(parent, bg=CARD, bd=0, highlightthickness=1, highlightbackground="#2A2A2A")
        return frame

    def _sync_main_scroll_region(self, _event: Any = None) -> None:
        if hasattr(self, "main_scroll_canvas") and self.main_scroll_canvas is not None:
            self.main_scroll_canvas.configure(scrollregion=self.main_scroll_canvas.bbox("all"))

    def _sync_main_scroll_width(self, event: Any) -> None:
        if hasattr(self, "main_scroll_canvas") and hasattr(self, "_main_scroll_window"):
            self.main_scroll_canvas.itemconfigure(self._main_scroll_window, width=event.width)

    def _on_main_mousewheel(self, event: Any) -> None:
        if not hasattr(self, "main_scroll_canvas") or self.main_scroll_canvas is None:
            return
        delta = 0
        if getattr(event, "delta", 0):
            delta = -1 * int(event.delta / 120)
        elif getattr(event, "num", None) == 4:
            delta = -1
        elif getattr(event, "num", None) == 5:
            delta = 1
        if delta:
            self.main_scroll_canvas.yview_scroll(delta, "units")

    def _build_layout(self) -> None:
        header = tk.Frame(self, bg=BG)
        header.pack(fill=tk.X, padx=12, pady=(10, 6))
        tk.Label(
            header,
            text="Endpoint Security Dashboard",
            bg=BG,
            fg=TEXT,
            font=("Segoe UI", 18, "bold"),
        ).pack(side=tk.LEFT)

        body = tk.Frame(self, bg=BG)
        body.pack(fill=tk.BOTH, expand=True, padx=12, pady=8)
        body.grid_rowconfigure(0, weight=1)
        body.grid_columnconfigure(1, weight=1)

        # Fixed left panel
        self.left_panel = self._card(body)
        self.left_panel.grid(row=0, column=0, sticky="ns", padx=(0, 8))
        self._build_left_panel()

        # Scrollable dashboard area
        scroll_shell = tk.Frame(body, bg=BG)
        scroll_shell.grid(row=0, column=1, sticky="nsew")
        scroll_shell.grid_rowconfigure(0, weight=1)
        scroll_shell.grid_columnconfigure(0, weight=1)

        self.main_scroll_canvas = tk.Canvas(scroll_shell, bg=BG, highlightthickness=0, borderwidth=0)
        main_scrollbar = ttk.Scrollbar(scroll_shell, orient="vertical", command=self.main_scroll_canvas.yview)
        self.main_scroll_canvas.configure(yscrollcommand=main_scrollbar.set)
        self.main_scroll_canvas.grid(row=0, column=0, sticky="nsew")
        main_scrollbar.grid(row=0, column=1, sticky="ns")

        self.scrollable_content = tk.Frame(self.main_scroll_canvas, bg=BG)
        self._main_scroll_window = self.main_scroll_canvas.create_window((0, 0), window=self.scrollable_content, anchor="nw")
        self.scrollable_content.bind("<Configure>", self._sync_main_scroll_region)
        self.main_scroll_canvas.bind("<Configure>", self._sync_main_scroll_width)
        self.main_scroll_canvas.bind_all("<MouseWheel>", self._on_main_mousewheel)
        self.main_scroll_canvas.bind_all("<Button-4>", self._on_main_mousewheel)
        self.main_scroll_canvas.bind_all("<Button-5>", self._on_main_mousewheel)

        self.right_panel = self._card(self.scrollable_content)
        self.right_panel.pack(fill=tk.X, expand=True, pady=(0, 12))
        self._build_right_panel()

        self.bottom_panel = self._card(self.scrollable_content)
        self.bottom_panel.pack(fill=tk.X, expand=True)
        self._build_bottom_panel()

    def _build_left_panel(self) -> None:
        layout = tk.Frame(self.left_panel, bg=CARD)
        layout.pack(fill=tk.BOTH, expand=True, padx=12, pady=10)
        layout.grid_columnconfigure(0, weight=1)
        layout.grid_rowconfigure(10, weight=1)

        tk.Label(
            layout,
            text="Systems",
            bg=CARD,
            fg=TEXT,
            font=("Segoe UI", 12, "bold"),
            anchor="w",
        ).grid(row=0, column=0, sticky="ew", pady=(0, 8))

        subnet_row = tk.Frame(layout, bg=CARD)
        subnet_row.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        subnet_row.grid_columnconfigure(1, weight=1)
        tk.Label(subnet_row, text="Subnet", bg=CARD, fg=MUTED, font=("Segoe UI", 9, "bold")).grid(
            row=0, column=0, sticky="w"
        )
        tk.Entry(
            subnet_row,
            textvariable=self.subnet_var,
            bg="#181818",
            fg=TEXT,
            insertbackground=TEXT,
            relief=tk.FLAT,
            width=18,
        ).grid(row=0, column=1, sticky="ew", padx=(8, 6))
        tk.Button(
            subnet_row,
            text="Discover",
            command=self._discover_devices_clicked,
            bg="#2C2C2C",
            fg=TEXT,
            relief=tk.FLAT,
        ).grid(row=0, column=2, sticky="e")

        self.discovery_status_label = tk.Label(
            layout,
            text="Discover devices to start remote scanning.",
            bg=CARD,
            fg=MUTED,
            anchor="w",
            justify=tk.LEFT,
            wraplength=320,
        )
        self.discovery_status_label.grid(row=2, column=0, sticky="ew", pady=(0, 8))

        self.remote_scan_progress = ttk.Progressbar(
            layout,
            orient=tk.HORIZONTAL,
            mode="determinate",
            maximum=100,
            style="Risk.Horizontal.TProgressbar",
        )
        self.remote_scan_progress.grid(row=3, column=0, sticky="ew", pady=(0, 4))

        self.scan_progress_label = tk.Label(
            layout,
            text="Scan idle.",
            bg=CARD,
            fg=MUTED,
            anchor="w",
            justify=tk.LEFT,
            wraplength=320,
        )
        self.scan_progress_label.grid(row=4, column=0, sticky="ew", pady=(0, 8))

        self.connection_status_label = tk.Label(
            layout,
            text="Not Connected",
            bg=CARD,
            fg=MUTED,
            anchor="w",
            justify=tk.LEFT,
            wraplength=320,
        )
        self.connection_status_label.grid(row=5, column=0, sticky="ew", pady=(0, 8))

        self.selected_device_label = tk.Label(
            layout,
            text="Select a device to view type and scan support.",
            bg=CARD,
            fg=MUTED,
            anchor="w",
            justify=tk.LEFT,
            wraplength=320,
        )
        self.selected_device_label.grid(row=6, column=0, sticky="ew", pady=(0, 10))

        btn_row_1 = tk.Frame(layout, bg=CARD)
        btn_row_1.grid(row=7, column=0, sticky="ew", pady=(0, 6))
        btn_row_1.grid_columnconfigure(0, weight=1)
        btn_row_1.grid_columnconfigure(1, weight=1)

        self.connect_scan_btn = tk.Button(
            btn_row_1,
            text="Connect & Scan",
            command=self._open_connect_dialog,
            bg="#2C2C2C",
            fg=TEXT,
            relief=tk.FLAT,
            state=tk.DISABLED,
        )
        self.connect_scan_btn.grid(row=0, column=0, sticky="ew", padx=(0, 4))

        self.refresh_scan_btn = tk.Button(
            btn_row_1,
            text="Reload / Refresh",
            command=self._refresh_selected_device_scan,
            bg="#2C2C2C",
            fg=TEXT,
            relief=tk.FLAT,
            state=tk.DISABLED,
        )
        self.refresh_scan_btn.grid(row=0, column=1, sticky="ew", padx=(4, 0))

        btn_row_2 = tk.Frame(layout, bg=CARD)
        btn_row_2.grid(row=8, column=0, sticky="ew", pady=(0, 8))
        btn_row_2.grid_columnconfigure(0, weight=1)
        btn_row_2.grid_columnconfigure(1, weight=1)
        btn_row_2.grid_columnconfigure(2, weight=1)

        self.local_scan_btn = tk.Button(
            btn_row_2,
            text="Local Scan",
            command=self._run_scan_clicked,
            bg="#2C2C2C",
            fg=TEXT,
            relief=tk.FLAT,
        )
        self.local_scan_btn.grid(row=0, column=0, sticky="ew", padx=(0, 4))

        self.unauth_scan_btn = tk.Button(
            btn_row_2,
            text="Scan Without Credentials",
            command=self._scan_without_credentials_clicked,
            bg="#2C2C2C",
            fg=TEXT,
            relief=tk.FLAT,
        )
        self.unauth_scan_btn.grid(row=0, column=1, sticky="ew", padx=4)

        tk.Button(
            btn_row_2,
            text="Word Report",
            command=self._generate_word_report,
            bg="#2C2C2C",
            fg=TEXT,
            relief=tk.FLAT,
        ).grid(row=0, column=2, sticky="ew", padx=(4, 0))

        btn_row_3 = tk.Frame(layout, bg=CARD)
        btn_row_3.grid(row=9, column=0, sticky="ew", pady=(0, 8))
        btn_row_3.grid_columnconfigure(0, weight=1)

        tk.Button(
            btn_row_3,
            text="Clear Saved Devices",
            command=self._clear_saved_devices,
            bg="#3A2424",
            fg=TEXT,
            relief=tk.FLAT,
        ).grid(row=0, column=0, sticky="ew")

        list_wrap = tk.Frame(layout, bg=CARD)
        list_wrap.grid(row=10, column=0, sticky="nsew")
        list_wrap.grid_rowconfigure(0, weight=1)
        list_wrap.grid_columnconfigure(0, weight=1)

        self.system_list = tk.Listbox(
            list_wrap,
            bg="#181818",
            fg=TEXT,
            selectbackground="#2E4A62",
            selectforeground=TEXT,
            borderwidth=0,
            highlightthickness=0,
            font=("Consolas", 10),
            exportselection=False,
            height=8,
        )
        system_scrollbar = ttk.Scrollbar(list_wrap, orient="vertical", command=self.system_list.yview)
        self.system_list.configure(yscrollcommand=system_scrollbar.set)
        self.system_list.grid(row=0, column=0, sticky="nsew")
        system_scrollbar.grid(row=0, column=1, sticky="ns")
        self.system_list.bind("<<ListboxSelect>>", self._on_system_select)

    def _format_eta_text(self, eta_seconds: int | None) -> str:
        if eta_seconds is None:
            return "estimating time remaining..."
        total_seconds = max(0, int(eta_seconds))
        minutes, seconds = divmod(total_seconds, 60)
        if minutes:
            return f"~{minutes}m {seconds:02d}s remaining"
        return f"~{seconds}s remaining"

    def _update_remote_scan_progress_ui(self, message: str, percent: float, eta_seconds: int | None = None) -> None:
        bounded_percent = max(0.0, min(100.0, float(percent)))
        self.remote_scan_progress["value"] = bounded_percent
        if message:
            self.discovery_status_label.config(text=message)

        if bounded_percent >= 100:
            detail = "100% complete | Results loaded."
        elif bounded_percent <= 0:
            detail = "Scan idle."
        else:
            detail = f"{int(round(bounded_percent))}% complete | {self._format_eta_text(eta_seconds)}"
        self.scan_progress_label.config(text=detail)

    def _reset_remote_scan_progress_ui(self, detail: str = "Scan idle.") -> None:
        self.remote_scan_progress["value"] = 0
        self.scan_progress_label.config(text=detail)

    def _any_scan_in_progress(self) -> bool:
        return bool(self.remote_scan_in_progress or self.unauth_scan_in_progress)

    def _get_connected_device_entry(self, ip_address: str) -> dict[str, Any] | None:
        entry = self.connected_devices.get(str(ip_address or "").strip())
        if not isinstance(entry, dict):
            return None
        if entry.get("session") is None:
            return None
        return entry

    def _device_type(self, device: dict[str, Any] | None) -> str:
        normalized = str((device or {}).get("device_type") or "Unknown").strip().title()
        if normalized not in {"Windows", "Mobile", "Unknown"}:
            return "Unknown"
        return normalized

    def _device_type_label(self, device: dict[str, Any] | None) -> str:
        return str((device or {}).get("device_type_label") or scanner.describe_device_type(self._device_type(device)))

    def _device_type_color(self, device_type: str) -> str:
        return DEVICE_TYPE_COLORS.get(str(device_type or "Unknown").title(), DEVICE_TYPE_COLORS["Unknown"])

    def _scan_supported_for_device(self, device: dict[str, Any] | None) -> bool:
        if not isinstance(device, dict):
            return False
        if "scan_supported" in device:
            return bool(device.get("scan_supported"))
        return self._device_type(device) != "Mobile"

    def _update_selected_device_details_ui(self) -> None:
        if not hasattr(self, "selected_device_label"):
            return

        device = self.selected_device
        if not isinstance(device, dict):
            self.selected_device_label.config(
                text="Select a device to view type and scan support.",
                fg=MUTED,
            )
            return

        ip_address = str(device.get("ip") or "").strip()
        primary_id = ip_address or self._device_display_name(device)
        status = str(device.get("status") or ("Online" if device.get("online") else "Offline")).strip() or "Unknown"
        device_type = self._device_type(device)
        type_label = self._device_type_label(device)
        is_local_device = self._is_local_device(device)
        has_auth_scan = isinstance(device.get("scan_result"), dict)
        has_unauth_scan = isinstance(device.get("unauth_scan_result"), dict)

        if is_local_device and not has_auth_scan:
            scan_status = "Local Scan Available"
        elif device_type == "Mobile":
            scan_status = "Not Supported"
        elif has_auth_scan and has_unauth_scan:
            scan_status = "Deep + External"
        elif has_auth_scan:
            scan_status = "Deep Scan Loaded"
        elif has_unauth_scan:
            scan_status = "Attacker View Loaded"
        elif not ip_address:
            scan_status = "Unavailable"
        elif "offline" in status.lower():
            scan_status = "Unavailable (Offline)"
        elif self._any_scan_in_progress():
            scan_status = "Busy"
        else:
            scan_status = "Supported"

        details = [
            f"Device: {primary_id}",
            f"Type: {type_label}",
            f"Status: {status}",
            f"Scan: {scan_status}",
        ]
        self.selected_device_label.config(
            text="\n".join(details),
            fg=self._device_type_color(device_type),
        )

    def _update_connection_status_ui(self) -> None:
        if not hasattr(self, "connection_status_label"):
            return

        device = self.selected_device
        ip_address = str((device or {}).get("ip") or "").strip()
        status_text = str((device or {}).get("status") or "").strip().lower()
        device_type = self._device_type(device)
        scan_supported = self._scan_supported_for_device(device)
        is_local_device = self._is_local_device(device)
        is_connected = bool(ip_address and self._get_connected_device_entry(ip_address))
        has_unauth_scan = bool(isinstance((device or {}).get("unauth_scan_result"), dict))

        if is_local_device and not (device or {}).get("scan_result"):
            self.connection_status_label.config(text="This device is ready for a local scan", fg="#66BB6A")
        elif device_type == "Mobile":
            self.connection_status_label.config(text="Mobile devices are not supported for scanning", fg="#FFB74D")
        elif is_connected:
            username = str(self.connected_devices[ip_address].get("username") or "").strip()
            username_suffix = f" as {username}" if username else ""
            self.connection_status_label.config(text=f"Connected ({ip_address}){username_suffix}", fg="#66BB6A")
        elif has_unauth_scan:
            self.connection_status_label.config(text="Attacker-view exposure data available (no credentials used)", fg="#FFB74D")
        else:
            self.connection_status_label.config(text="Not Connected", fg=MUTED)

        is_offline = "offline" in status_text or "failed" in status_text
        busy = self._any_scan_in_progress()
        can_connect = bool(ip_address) and not is_local_device and not is_offline and scan_supported and not busy
        can_refresh = bool(ip_address) and is_connected and scan_supported and not busy
        if hasattr(self, "connect_scan_btn"):
            self.connect_scan_btn.config(state=tk.NORMAL if can_connect else tk.DISABLED)
        if hasattr(self, "refresh_scan_btn"):
            self.refresh_scan_btn.config(state=tk.NORMAL if can_refresh else tk.DISABLED)
        if hasattr(self, "local_scan_btn"):
            self.local_scan_btn.config(state=tk.DISABLED if busy else tk.NORMAL)
        if hasattr(self, "unauth_scan_btn"):
            self.unauth_scan_btn.config(state=tk.DISABLED if busy else tk.NORMAL)
        self._update_selected_device_details_ui()

    def _device_key(self, device: dict[str, Any]) -> str:
        return str(device.get("device_id") or device.get("ip") or device.get("hostname") or "Unknown")

    def _device_display_name(self, device: dict[str, Any]) -> str:
        hostname = str(device.get("hostname") or "").strip()
        ip_address = str(device.get("ip") or "").strip()
        if hostname and ip_address and hostname.lower() != ip_address.lower():
            return f"{hostname} ({ip_address})"
        return hostname or ip_address or "Unknown"

    def _device_from_scan(self, scan: dict[str, Any]) -> dict[str, Any]:
        enriched_scan = scanner.enrich_device_metadata(scan)
        system_info = enriched_scan.get("system_info") or {}
        ip_address = str(enriched_scan.get("target_ip") or system_info.get("ip_address") or "").strip()
        hostname = str(enriched_scan.get("hostname") or system_info.get("hostname") or ip_address or "Unknown").strip()
        status = str(enriched_scan.get("connection_status") or "Scanned").strip() or "Scanned"
        return {
            "device_id": scanner.get_device_id(enriched_scan),
            "ip": ip_address,
            "hostname": hostname,
            "online": status not in {"Offline", "Connection Failed", "Connection Failed (Timeout)"},
            "status": status,
            "last_scan_at": enriched_scan.get("last_scan_at"),
            "device_type": str(enriched_scan.get("device_type") or "Unknown"),
            "device_type_label": str(enriched_scan.get("device_type_label") or scanner.describe_device_type(enriched_scan.get("device_type"))),
            "scan_supported": bool(enriched_scan.get("scan_supported", True)),
            "device_type_reason": str(enriched_scan.get("device_type_reason") or ""),
            "mac_address": str(enriched_scan.get("mac_address") or ""),
            "mac_vendor": str(enriched_scan.get("mac_vendor") or ""),
            "scan_result": enriched_scan,
            "unauth_scan_result": None,
        }

    def _device_from_unauth_scan(self, scan: dict[str, Any]) -> dict[str, Any]:
        enriched_scan = scanner.enrich_device_metadata(scan)
        system_info = enriched_scan.get("system_info") or {}
        ip_address = str(enriched_scan.get("target_ip") or system_info.get("ip_address") or "").strip()
        hostname = str(enriched_scan.get("hostname") or system_info.get("hostname") or ip_address or "Unknown").strip()
        exposure = enriched_scan.get("external_exposure") or {}
        return {
            "device_id": scanner.get_device_id(enriched_scan),
            "ip": ip_address,
            "hostname": hostname,
            "online": True,
            "status": "Exposure Assessed",
            "last_scan_at": enriched_scan.get("last_scan_at"),
            "device_type": str(enriched_scan.get("device_type") or "Unknown"),
            "device_type_label": str(enriched_scan.get("device_type_label") or scanner.describe_device_type(enriched_scan.get("device_type"))),
            "scan_supported": bool(enriched_scan.get("scan_supported", True)),
            "device_type_reason": str(enriched_scan.get("device_type_reason") or ""),
            "mac_address": str(exposure.get("mac_address") or enriched_scan.get("mac_address") or ""),
            "mac_vendor": str(exposure.get("mac_vendor") or enriched_scan.get("mac_vendor") or ""),
            "scan_result": None,
            "unauth_scan_result": enriched_scan,
        }

    def _is_local_device(self, device: dict[str, Any] | None) -> bool:
        return bool(isinstance(device, dict) and device.get("is_local_device"))

    def _make_local_device_placeholder(self) -> dict[str, Any]:
        try:
            system_info = scanner.os_check.get_system_info()
        except Exception:
            system_info = {}

        hostname = str(system_info.get("hostname") or "This Device").strip() or "This Device"
        ip_address = str(system_info.get("ip_address") or "").strip()
        placeholder_scan = scanner.enrich_device_metadata(
            {
                "hostname": hostname,
                "device_id": ip_address or hostname,
                "target_ip": ip_address,
                "connection_status": "Local Device",
                "system_info": system_info,
            }
        )
        return {
            "device_id": scanner.get_device_id(placeholder_scan),
            "ip": ip_address,
            "hostname": hostname,
            "online": True,
            "status": "Local Device",
            "last_scan_at": None,
            "device_type": str(placeholder_scan.get("device_type") or "Windows"),
            "device_type_label": str(
                placeholder_scan.get("device_type_label")
                or scanner.describe_device_type(placeholder_scan.get("device_type"))
            ),
            "scan_supported": bool(placeholder_scan.get("scan_supported", True)),
            "device_type_reason": str(placeholder_scan.get("device_type_reason") or ""),
            "mac_address": str(placeholder_scan.get("mac_address") or ""),
            "mac_vendor": str(placeholder_scan.get("mac_vendor") or ""),
            "scan_result": None,
            "unauth_scan_result": None,
            "is_local_device": True,
        }

    def _upsert_device(self, device: dict[str, Any]) -> None:
        key = self._device_key(device).lower()
        for idx, current in enumerate(self.devices):
            if self._device_key(current).lower() == key:
                merged = dict(current)
                merged.update({k: v for k, v in device.items() if v is not None and v != ""})
                if device.get("scan_result") is None and current.get("scan_result") is not None:
                    merged["scan_result"] = current["scan_result"]
                if device.get("unauth_scan_result") is None and current.get("unauth_scan_result") is not None:
                    merged["unauth_scan_result"] = current["unauth_scan_result"]
                self.devices[idx] = merged
                return
        self.devices.append(device)

    def _sync_devices_from_results(self) -> None:
        for scan in self.results:
            self._upsert_device(self._device_from_scan(scan))
        for scan in self.unauth_results:
            self._upsert_device(self._device_from_unauth_scan(scan))
        self.devices.sort(
            key=lambda item: (
                0 if item.get("online") else 1,
                self._device_display_name(item).lower(),
            )
        )

    def _clear_saved_scan_snapshots(self) -> None:
        for device in self.devices:
            if not isinstance(device, dict):
                continue
            device["scan_result"] = None
            device["unauth_scan_result"] = None
            if str(device.get("status") or "").strip() == "Exposure Assessed":
                device["status"] = "Online" if device.get("online") else "Offline"

    def _select_device_by_key(self, device_key: str) -> None:
        target = str(device_key or "").lower()
        for idx, device in enumerate(self.device_rows):
            if self._device_key(device).lower() == target:
                self.system_list.selection_clear(0, tk.END)
                self.system_list.selection_set(idx)
                self.system_list.activate(idx)
                self.system_list.see(idx)
                self._apply_device_selection(idx)
                return

    def _make_placeholder_scan(self, device: dict[str, Any]) -> dict[str, Any]:
        hostname = str(device.get("hostname") or device.get("ip") or "Unknown")
        ip_address = str(device.get("ip") or "")
        placeholder = {
            "hostname": hostname,
            "device_id": self._device_key(device),
            "target_ip": ip_address,
            "risk_score": 0,
            "risk_level": "UNKNOWN",
            "connection_status": str(device.get("status") or "Not Scanned"),
            "vulnerabilities": [],
            "system_info": {"hostname": hostname, "ip_address": ip_address},
            "antivirus": {},
            "patches": {},
            "firewall": {"profiles": {}},
            "open_ports": {"listening_ports": [], "risky_open": []},
            "services": {"highlighted": []},
            "third_party_software": {"risky_apps": []},
            "installed_apps": [],
            "risky_apps": [],
            "device_type": self._device_type(device),
            "device_type_label": self._device_type_label(device),
            "scan_supported": self._scan_supported_for_device(device),
        }
        return scanner.enrich_device_metadata(placeholder)

    def _apply_device_selection(self, index: int) -> None:
        if not (0 <= index < len(self.device_rows)):
            return
        device = self.device_rows[index]
        self.selected_device = device
        self.selected_unauth_scan = device.get("unauth_scan_result") if isinstance(device.get("unauth_scan_result"), dict) else None
        self._update_connection_status_ui()

        scan = device.get("scan_result")
        if isinstance(scan, dict):
            self.selected_scan = scan
            self._render_selected_scan()
            return

        unauth_scan = device.get("unauth_scan_result")
        if isinstance(unauth_scan, dict):
            self.selected_scan = unauth_scan
            self._render_selected_scan()
            return

        self.selected_scan = self._make_placeholder_scan(device)
        self._render_selected_scan()
        if self._is_local_device(device):
            self.why_label.config(
                text=f"{self._device_display_name(device)} is this computer. Click Local Scan to capture fresh results."
            )
            self.top_drivers_label.config(text="Top Risk Drivers: Run Local Scan to load current data.", fg=MUTED)
            return
        if self._device_type(device) == "Mobile":
            self.why_label.config(
                text=f"{self._device_display_name(device)} is a mobile device. Mobile devices are not supported for WinRM scanning."
            )
            self.top_drivers_label.config(text="Top Risk Drivers: Scan not supported for mobile devices.", fg="#FFB74D")
        else:
            self.why_label.config(
                text=(
                    f"{self._device_display_name(device)} is {device.get('status', 'Not Scanned')}. "
                    "Click Connect & Scan for a deep scan, or Scan Without Credentials for attacker-view exposure data."
                )
            )
            self.top_drivers_label.config(text="Top Risk Drivers: No scan data loaded for this device.", fg=MUTED)

    def _discover_devices_clicked(self) -> None:
        subnet = str(self.subnet_var.get() or "").strip()
        if not subnet:
            messagebox.showwarning("Discover Devices", "Enter a subnet such as 192.168.1.0/24 first.")
            return

        self.discovery_status_label.config(text=f"Discovering devices in {subnet}...")
        threading.Thread(target=self._discover_devices_worker, args=(subnet,), daemon=True).start()

    def _discover_devices_worker(self, subnet: str) -> None:
        try:
            devices = scanner.discover_devices(subnet)
            self.after(0, lambda subnet=subnet, devices=devices: self._on_discovery_complete(subnet, devices))
        except Exception as exc:
            message = str(exc)
            self.after(0, lambda message=message: self._on_discovery_failed(message))

    def _on_discovery_complete(self, subnet: str, devices: list[dict[str, Any]]) -> None:
        for device in devices:
            self._upsert_device(
                {
                    "device_id": device.get("ip") or device.get("hostname") or "Unknown",
                    "ip": device.get("ip") or "",
                    "hostname": device.get("hostname") or "",
                    "online": bool(device.get("online")),
                    "status": str(device.get("status") or ("Online" if device.get("online") else "Offline")),
                    "last_scan_at": device.get("last_scan_at"),
                    "device_type": device.get("device_type") or "Unknown",
                    "device_type_label": device.get("device_type_label") or scanner.describe_device_type(device.get("device_type")),
                    "scan_supported": bool(device.get("scan_supported", True)),
                    "device_type_reason": device.get("device_type_reason") or "",
                    "mac_address": device.get("mac_address") or "",
                    "mac_vendor": device.get("mac_vendor") or "",
                    "scan_result": None,
                }
            )
        self._sync_devices_from_results()
        self._populate_system_list()
        online_count = sum(1 for item in devices if item.get("online"))
        self.discovery_status_label.config(
            text=f"Discovery complete for {subnet}: {online_count} online, {max(0, len(devices) - online_count)} offline."
        )
        if self.device_rows and self.selected_device is None:
            self._select_device_by_key(self._device_key(self.device_rows[0]))

    def _on_discovery_failed(self, message: str) -> None:
        self.discovery_status_label.config(text=f"Discovery failed: {message}")
        messagebox.showerror("Discovery Failed", message)

    def _scan_without_credentials_clicked(self) -> None:
        subnet = str(self.subnet_var.get() or "").strip()
        if not subnet:
            messagebox.showwarning("Scan Without Credentials", "Enter a subnet such as 192.168.1.0/24 first.")
            return
        if self._any_scan_in_progress():
            messagebox.showinfo("Scan Without Credentials", "Wait for the current scan to finish first.")
            return

        self.unauth_scan_in_progress = True
        self._update_remote_scan_progress_ui(
            f"Running unauthenticated exposure scan on {subnet}...",
            2,
            None,
        )
        self.discovery_status_label.config(
            text=f"Running attacker-view network reconnaissance on {subnet} without credentials..."
        )
        self._update_connection_status_ui()
        threading.Thread(target=self._unauthenticated_scan_worker, args=(subnet,), daemon=True).start()

    def _unauthenticated_scan_worker(self, subnet: str) -> None:
        try:
            def _progress(message: str, percent: float, eta_seconds: int | None) -> None:
                self.after(
                    0,
                    lambda message=message, percent=percent, eta_seconds=eta_seconds: self._update_remote_scan_progress_ui(
                        message,
                        percent,
                        eta_seconds,
                    ),
                )

            results = scanner.add_unauthenticated_scan_to_results(subnet, progress_callback=_progress)
            self.after(0, lambda subnet=subnet, results=results: self._on_unauthenticated_scan_success(subnet, results))
        except Exception as exc:
            message = str(exc)
            self.after(0, lambda message=message: self._on_unauthenticated_scan_failure(message))

    def _on_unauthenticated_scan_success(self, subnet: str, results: list[dict[str, Any]]) -> None:
        self.unauth_scan_in_progress = False
        self.unauth_results = scanner.load_unauthenticated_scan_results()
        self._refresh_all()
        online_count = len(results)
        self._update_remote_scan_progress_ui(
            f"Unauthenticated scan completed for {subnet}.",
            100,
            0,
        )
        self.discovery_status_label.config(
            text=f"Unauthenticated scan completed for {subnet}: {online_count} active device(s) assessed."
        )
        if results:
            self._select_device_by_key(scanner.get_device_id(results[0]))
        self._update_connection_status_ui()

    def _on_unauthenticated_scan_failure(self, message: str) -> None:
        self.unauth_scan_in_progress = False
        self._reset_remote_scan_progress_ui("Unauthenticated scan failed.")
        self.discovery_status_label.config(text=f"Unauthenticated scan failed: {message}")
        self._update_connection_status_ui()
        messagebox.showerror("Scan Without Credentials Failed", message)

    def _refresh_selected_device_scan(self) -> None:
        if not self.selected_device or not self.selected_device.get("ip"):
            messagebox.showwarning("Reload / Refresh", "Select a connected device with an IP address first.")
            return
        if not self._scan_supported_for_device(self.selected_device):
            messagebox.showinfo("Reload / Refresh", "Mobile devices are not supported for scanning.")
            return
        ip_address = str(self.selected_device.get("ip") or "").strip()
        if not self._get_connected_device_entry(ip_address):
            messagebox.showinfo("Reload / Refresh", "This device is not connected. Use Connect & Scan first.")
            return
        self._start_remote_scan(ip_address, reuse_cached_session=True)

    def _open_connect_dialog(self) -> None:
        if not self.selected_device or not self.selected_device.get("ip"):
            messagebox.showwarning("Connect & Scan", "Select a discovered device with an IP address first.")
            return
        if not self._scan_supported_for_device(self.selected_device):
            messagebox.showinfo("Connect & Scan", "Mobile devices are not supported for scanning.")
            return

        ip_address = str(self.selected_device.get("ip") or "").strip()
        if self._get_connected_device_entry(ip_address):
            self._start_remote_scan(ip_address, reuse_cached_session=True)
            return

        dialog = tk.Toplevel(self)
        dialog.title("Connect & Scan")
        dialog.configure(bg=CARD)
        dialog.transient(self)
        dialog.grab_set()
        dialog.resizable(False, False)

        ip_var = tk.StringVar(value=str(self.selected_device.get("ip") or ""))
        user_var = tk.StringVar()
        password_var = tk.StringVar()

        fields = [
            ("IP Address", ip_var, True),
            ("Username", user_var, False),
            ("Password", password_var, False),
        ]
        for row_index, (label_text, variable, readonly) in enumerate(fields):
            tk.Label(dialog, text=label_text, bg=CARD, fg=TEXT, anchor="w").grid(
                row=row_index, column=0, sticky="w", padx=12, pady=(12 if row_index == 0 else 8, 0)
            )
            entry = tk.Entry(
                dialog,
                textvariable=variable,
                bg="#181818",
                fg=TEXT,
                insertbackground=TEXT,
                relief=tk.FLAT,
                width=36,
                show="*" if label_text == "Password" else "",
            )
            if readonly:
                entry.config(state="readonly", readonlybackground="#181818")
            entry.grid(row=row_index, column=1, sticky="ew", padx=12, pady=(12 if row_index == 0 else 8, 0))

        tk.Label(
            dialog,
            text="Plain local usernames are auto-tried as .\\username. If needed, use COMPUTERNAME\\username.",
            bg=CARD,
            fg=MUTED,
            anchor="w",
            justify=tk.LEFT,
        ).grid(row=len(fields), column=0, columnspan=2, sticky="w", padx=12, pady=(8, 0))

        dialog.grid_columnconfigure(1, weight=1)

        btn_row = tk.Frame(dialog, bg=CARD)
        btn_row.grid(row=len(fields) + 1, column=0, columnspan=2, sticky="ew", padx=12, pady=12)

        def _connect() -> None:
            username = str(user_var.get() or "").strip()
            password = password_var.get()
            if not username or not password:
                messagebox.showwarning("Connect & Scan", "Enter both username and password.", parent=dialog)
                return
            selected_ip = str(ip_var.get() or "").strip()
            password_var.set("")
            dialog.grab_release()
            dialog.destroy()
            self._start_remote_scan(selected_ip, username=username, password=password, reuse_cached_session=False)

        tk.Button(btn_row, text="Connect", command=_connect, bg="#2C2C2C", fg=TEXT, relief=tk.FLAT).pack(
            side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 6)
        )
        tk.Button(btn_row, text="Cancel", command=dialog.destroy, bg="#2C2C2C", fg=TEXT, relief=tk.FLAT).pack(
            side=tk.LEFT, fill=tk.X, expand=True
        )

    def _start_remote_scan(
        self,
        ip_address: str,
        username: str | None = None,
        password: str | None = None,
        reuse_cached_session: bool = False,
    ) -> None:
        if self.selected_device and not self._scan_supported_for_device(self.selected_device):
            self._update_connection_status_ui()
            return
        self.remote_scan_in_progress = True
        self._update_remote_scan_progress_ui(
            f"Connecting to {ip_address} and preparing the remote scan...",
            2,
            60,
        )
        self._update_connection_status_ui()
        threading.Thread(
            target=self._remote_scan_worker,
            args=(ip_address, username, password, reuse_cached_session),
            daemon=True,
        ).start()

    def _remote_scan_worker(
        self,
        ip_address: str,
        username: str | None,
        password: str | None,
        reuse_cached_session: bool,
    ) -> None:
        try:
            connection_entry = self._get_connected_device_entry(ip_address) if reuse_cached_session else None
            if connection_entry:
                session = connection_entry["session"]
                username_for_cache = str(connection_entry.get("username") or "").strip()
            else:
                session = scanner.create_winrm_session(ip_address, str(username or ""), str(password or ""))
                username_for_cache = str(
                    getattr(session, "_endpoint_dashboard_username", str(username or "").strip()) or ""
                ).strip()

            def _progress(message: str, percent: float, eta_seconds: int | None) -> None:
                self.after(
                    0,
                    lambda message=message, percent=percent, eta_seconds=eta_seconds: self._update_remote_scan_progress_ui(
                        message,
                        percent,
                        eta_seconds,
                    ),
                )

            scan = scanner.run_remote_scan_with_session(
                ip_address,
                session,
                progress_callback=_progress,
            )
            self.after(
                0,
                lambda scan=scan, session=session, ip_address=ip_address, username_for_cache=username_for_cache: self._on_remote_scan_success(
                    scan,
                    ip_address,
                    session,
                    username_for_cache,
                ),
            )
        except scanner.RemoteScanError as exc:
            self.after(
                0,
                lambda ip_address=ip_address, message=str(exc), code=exc.code, reuse_cached_session=reuse_cached_session: self._on_remote_scan_failure(
                    ip_address,
                    scanner.RemoteScanError(message, code=code),
                    reuse_cached_session,
                ),
            )
        except Exception as exc:
            wrapped_message = str(exc)
            self.after(
                0,
                lambda ip_address=ip_address, wrapped_message=wrapped_message, reuse_cached_session=reuse_cached_session: self._on_remote_scan_failure(
                    ip_address,
                    scanner.RemoteScanError(wrapped_message, code="execution_error"),
                    reuse_cached_session,
                ),
            )

    def _on_remote_scan_success(self, scan: dict[str, Any], ip_address: str, session: Any, username: str) -> None:
        self.remote_scan_in_progress = False
        self.connected_devices[ip_address] = {
            "session": session,
            "username": username,
            "connected_at": datetime.now().isoformat(),
        }
        scanner.upsert_scan_result(self.results, scan)
        scanner.save_scan_results(self.results)
        self._upsert_device(self._device_from_scan(scan))
        self._sync_devices_from_results()
        self._populate_system_list()
        self._update_remote_scan_progress_ui(
            f"Remote scan completed for {scan.get('target_ip') or scan.get('hostname')}.",
            100,
            0,
        )
        self.discovery_status_label.config(text=f"Remote scan completed for {scan.get('target_ip') or scan.get('hostname')}.")
        self._select_device_by_key(scanner.get_device_id(scan))
        self._update_connection_status_ui()

    def _on_remote_scan_failure(
        self,
        ip_address: str,
        error: scanner.RemoteScanError,
        reuse_cached_session: bool = False,
    ) -> None:
        self.remote_scan_in_progress = False
        display_error = error
        if reuse_cached_session and error.code in {"invalid_credentials", "connection_failed", "timeout"}:
            self.connected_devices.pop(ip_address, None)
            display_error = scanner.RemoteScanError("Session expired, please reconnect.", code="session_expired")
        online = error.code not in {"connection_failed", "timeout"}
        self._upsert_device(
            {
                "device_id": ip_address,
                "ip": ip_address,
                "hostname": "",
                "online": online,
                "status": str(display_error),
                "last_scan_at": None,
                "scan_result": None,
            }
        )
        self._sync_devices_from_results()
        self._populate_system_list()
        self._reset_remote_scan_progress_ui("Remote scan failed.")
        self.discovery_status_label.config(text=f"{ip_address}: {display_error}")
        self._select_device_by_key(ip_address)
        self._update_connection_status_ui()
        messagebox.showerror("Remote Scan Failed", str(display_error))

    def _build_right_panel(self) -> None:
        self.right_view_mode = "risk"

        # View switching buttons
        switch_frame = tk.Frame(self.right_panel, bg=CARD)
        switch_frame.pack(fill=tk.X, padx=12, pady=(10, 6))

        active_bg = "#2C2C2C"
        inactive_bg = "#1E1E1E"

        self.risk_btn = tk.Button(
            switch_frame,
            text="Risk Analysis",
            bg=active_bg if self.right_view_mode in ("risk", "split") else inactive_bg,
            fg=TEXT,
            relief=tk.FLAT,
            command=lambda: self._switch_right_view("risk"),
        )
        self.risk_btn.pack(side=tk.LEFT, padx=(0, 6), fill=tk.X, expand=True)

        self.threat_btn = tk.Button(
            switch_frame,
            text="Threat Analysis",
            bg=active_bg if self.right_view_mode in ("threat", "split") else inactive_bg,
            fg=TEXT,
            relief=tk.FLAT,
            command=lambda: self._switch_right_view("threat"),
        )
        self.threat_btn.pack(side=tk.LEFT, padx=6, fill=tk.X, expand=True)

        self.split_btn = tk.Button(
            switch_frame,
            text="Split View",
            bg=active_bg if self.right_view_mode == "split" else inactive_bg,
            fg=TEXT,
            relief=tk.FLAT,
            command=lambda: self._switch_right_view("split"),
        )
        self.split_btn.pack(side=tk.LEFT, padx=(6, 0), fill=tk.X, expand=True)

        # Resizable split view container
        self.split_paned = tk.PanedWindow(
            self.right_panel,
            orient=tk.HORIZONTAL,
            sashrelief=tk.RAISED,
            sashwidth=6,
        )
        self.split_paned.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 6))

        # Risk (charts) container (left pane)
        self.risk_container = tk.Frame(self.split_paned, bg=CARD)

        tk.Label(
            self.risk_container,
            text="Risk Contribution Analysis",
            bg=CARD,
            fg=TEXT,
            font=("Segoe UI", 12, "bold"),
        ).pack(anchor="w", padx=12, pady=(10, 4))

        chart_wrap = tk.Frame(self.risk_container, bg=CARD)
        chart_wrap.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 6))
        self.figure = Figure(figsize=(6, 3.8), dpi=100, facecolor=CARD)
        self.ax_pie = self.figure.add_subplot(121)
        self.ax_bar = self.figure.add_subplot(122)
        self.ax_pie.set_facecolor(CARD)
        self.ax_bar.set_facecolor(CARD)
        self.canvas = FigureCanvasTkAgg(self.figure, master=chart_wrap)
        self.canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self.canvas.mpl_connect("button_press_event", self._on_pie_click)

        self.summary_label = tk.Label(
            self.risk_container, text="", bg=CARD, fg=MUTED, justify=tk.LEFT, anchor="w"
        )
        self.summary_label.pack(fill=tk.X, padx=12, pady=(0, 4))

        self.pie_click_info = tk.Label(
            self.risk_container, text="", bg=CARD, fg=ACCENT, justify=tk.LEFT, anchor="w"
        )
        self.pie_click_info.pack(fill=tk.X, padx=12, pady=(0, 8))

        self.top_drivers_label = tk.Label(
            self.risk_container, text="", bg=CARD, fg=TEXT, justify=tk.LEFT, anchor="w"
        )
        self.top_drivers_label.pack(fill=tk.X, padx=12, pady=(0, 8))

        contrib_card = tk.Frame(self.risk_container, bg="#171717")
        contrib_card.pack(fill=tk.X, padx=12, pady=(0, 12))
        tk.Label(
            contrib_card,
            text="Component Risk Breakdown",
            bg="#171717",
            fg=TEXT,
            font=("Segoe UI", 10, "bold"),
        ).pack(anchor="w", padx=10, pady=(8, 8))
        for comp in risk.COMPONENTS:
            row = tk.Frame(contrib_card, bg="#171717")
            row.pack(fill=tk.X, padx=10, pady=3)
            tk.Label(row, text=f"{comp:<12}", bg="#171717", fg=TEXT, width=14, anchor="w").pack(side=tk.LEFT)
            pb = ttk.Progressbar(
                row,
                style="Risk.Horizontal.TProgressbar",
                orient=tk.HORIZONTAL,
                mode="determinate",
                length=220,
            )
            pb.pack(side=tk.LEFT, padx=6, fill=tk.X, expand=True)
            lbl = tk.Label(row, text="0 pts (0%)", bg="#171717", fg=MUTED, width=16, anchor="e")
            lbl.pack(side=tk.RIGHT)
            self.component_bars[comp] = pb
            self.component_labels[comp] = lbl

        # Threat container (right pane)
        self.threat_container = tk.Frame(self.split_paned, bg=CARD)

        tk.Label(
            self.threat_container,
            text="Threat Analysis Dashboard",
            bg=CARD,
            fg=TEXT,
            font=("Segoe UI", 14, "bold"),
        ).pack(anchor="w", padx=12, pady=(10, 4))

        tk.Label(
            self.threat_container,
            text="Compliance: NIST SP 800-30 aligned risk scoring & NIST CSF 2.0 mapping",
            bg=CARD,
            fg=MUTED,
            justify=tk.LEFT,
            anchor="w",
            font=("Segoe UI", 10, "bold"),
        ).pack(fill=tk.X, padx=12, pady=(0, 10))

        export_row = tk.Frame(self.threat_container, bg=CARD)
        export_row.pack(fill=tk.X, padx=12, pady=(0, 6))
        tk.Button(
            export_row,
            text="Export Threat Report",
            bg="#2C2C2C",
            fg=TEXT,
            relief=tk.FLAT,
            command=self._export_threat_report_clicked,
        ).pack(side=tk.LEFT)
        tk.Button(
            export_row,
            text="Export to Excel",
            bg="#2C2C2C",
            fg=TEXT,
            relief=tk.FLAT,
            command=self._export_threats_to_excel_clicked,
        ).pack(side=tk.RIGHT)

        # Threat charts (new, count-based)
        charts_wrap = tk.Frame(self.threat_container, bg=CARD)
        charts_wrap.pack(fill=tk.X, padx=12, pady=(0, 10))

        self.threat_figure = Figure(figsize=(6, 3.1), dpi=100, facecolor=CARD)
        self.threat_ax_sev = self.threat_figure.add_subplot(121)
        self.threat_ax_cat = self.threat_figure.add_subplot(122)
        self.threat_ax_sev.set_facecolor(CARD)
        self.threat_ax_cat.set_facecolor(CARD)
        self.threat_canvas = FigureCanvasTkAgg(self.threat_figure, master=charts_wrap)
        self.threat_canvas.get_tk_widget().pack(fill=tk.X, expand=True)

        self.threat_figure.tight_layout(pad=1.0)
        self.threat_canvas.draw_idle()

        # Scrollable area for threat cards
        threat_wrap = tk.Frame(self.threat_container, bg=CARD)
        threat_wrap.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 6))

        canvas = tk.Canvas(threat_wrap, bg=CARD, highlightthickness=0)
        vscroll = tk.Scrollbar(threat_wrap, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vscroll.set)
        vscroll.pack(side=tk.RIGHT, fill=tk.Y)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        inner = tk.Frame(canvas, bg=CARD)
        inner_id = canvas.create_window((0, 0), window=inner, anchor="nw")

        def _on_frame_config(_event: Any = None) -> None:
            canvas.configure(scrollregion=canvas.bbox("all"))
            try:
                canvas.itemconfigure(inner_id, width=canvas.winfo_width())
            except Exception:
                pass

        inner.bind("<Configure>", _on_frame_config)
        canvas.bind("<Configure>", _on_frame_config)
        self.threat_inner_frame = inner

        # Initialize with risk-only layout
        self._switch_right_view(self.right_view_mode)

    def _switch_right_view(self, mode: str) -> None:
        self.right_view_mode = mode

        active_bg = "#2C2C2C"
        inactive_bg = "#1E1E1E"

        # Toggle button highlight
        try:
            self.risk_btn.config(bg=active_bg if mode in ("risk", "split") else inactive_bg)
            self.threat_btn.config(bg=active_bg if mode in ("threat", "split") else inactive_bg)
            self.split_btn.config(bg=active_bg if mode == "split" else inactive_bg)
        except Exception:
            pass

        # Replace panes depending on view mode
        try:
            for pane in list(self.split_paned.panes()):
                self.split_paned.forget(pane)
        except Exception:
            pass

        if mode in ("risk", "split"):
            self.split_paned.add(self.risk_container, minsize=320, stretch="always")
        if mode in ("threat", "split"):
            self.split_paned.add(self.threat_container, minsize=380, stretch="always")

        if mode == "split":
            total = self.split_paned.winfo_width()
            if total <= 1:
                total = 1000
            self.split_paned.sash_place(0, int(total * 0.52), 0)

        if mode in ("threat", "split") and self.selected_scan:
            self._render_threat_analysis(self.selected_scan)

    def _build_bottom_panel(self) -> None:
        tk.Label(self.bottom_panel, text="Detailed Security Components", bg=CARD, fg=TEXT, font=("Segoe UI", 12, "bold")).pack(
            anchor="w", padx=12, pady=(10, 2)
        )
        self.why_label = tk.Label(self.bottom_panel, text="", bg=CARD, fg=MUTED, anchor="w", justify=tk.LEFT)
        self.why_label.pack(fill=tk.X, padx=12, pady=(0, 8))

        self.notebook = ttk.Notebook(self.bottom_panel, style="Dark.TNotebook")
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 10))

        self.tab_antivirus = self._make_text_tab("Antivirus")
        self.tab_firewall = self._make_text_tab("Firewall")
        self.tab_ports = self._make_text_tab("Open Ports")
        self._make_external_exposure_tab("External Exposure")
        self.tab_patching = self._make_text_tab("Patching")
        self._make_software_inventory_tab("Third-party Software")
        self._make_live_applications_tab("Live Applications")
        self.tab_raw = self._make_text_tab("Raw JSON")
        self.notebook.bind("<<NotebookTabChanged>>", self._on_detail_tab_changed)

    def _make_text_tab(self, title: str) -> tk.Text:
        frame = tk.Frame(self.notebook, bg=CARD)
        self.notebook.add(frame, text=title)
        txt = tk.Text(
            frame,
            bg="#181818",
            fg=TEXT,
            insertbackground=TEXT,
            relief=tk.FLAT,
            borderwidth=0,
            wrap=tk.WORD,
            padx=8,
            pady=8,
            font=("Consolas", 10),
        )
        txt.pack(fill=tk.BOTH, expand=True)
        return txt

    def _make_external_exposure_tab(self, title: str) -> None:
        frame = tk.Frame(self.notebook, bg=CARD)
        self.notebook.add(frame, text=title)

        header = tk.Frame(frame, bg=CARD)
        header.pack(fill=tk.X, padx=12, pady=(10, 6))

        self.external_exposure_title_label = tk.Label(
            header,
            text="Attacker View: No data loaded",
            bg=CARD,
            fg=TEXT,
            font=("Segoe UI", 12, "bold"),
            anchor="w",
            justify=tk.LEFT,
        )
        self.external_exposure_title_label.pack(anchor="w", fill=tk.X)

        self.external_exposure_note_label = tk.Label(
            header,
            text="Run Scan Without Credentials to populate attacker-view exposure data.",
            bg=CARD,
            fg="#FFB74D",
            anchor="w",
            justify=tk.LEFT,
            wraplength=980,
        )
        self.external_exposure_note_label.pack(anchor="w", fill=tk.X, pady=(6, 0))

        grid = tk.Frame(frame, bg=CARD)
        grid.pack(fill=tk.X, padx=12, pady=(0, 8))
        grid.grid_columnconfigure(1, weight=1)
        grid.grid_columnconfigure(3, weight=1)

        exposure_fields = [
            ("IP Address", "ip_address"),
            ("Hostname", "hostname"),
            ("Device Type", "device_type"),
            ("MAC / Vendor", "mac_vendor"),
            ("Risk Level", "risk_level"),
            ("Open Ports", "open_ports"),
            ("Scan Type", "scan_type"),
            ("Visibility", "visibility"),
        ]
        self.external_exposure_value_labels: dict[str, tk.Label] = {}
        for index, (label_text, field_key) in enumerate(exposure_fields):
            row = index // 2
            column = (index % 2) * 2
            tk.Label(
                grid,
                text=label_text,
                bg=CARD,
                fg=MUTED,
                font=("Segoe UI", 9, "bold"),
                anchor="w",
            ).grid(row=row, column=column, sticky="w", padx=(0, 8), pady=(4, 0))
            value_label = tk.Label(
                grid,
                text="-",
                bg=CARD,
                fg=TEXT,
                anchor="w",
                justify=tk.LEFT,
                wraplength=420,
            )
            value_label.grid(row=row, column=column + 1, sticky="ew", padx=(0, 16), pady=(4, 0))
            self.external_exposure_value_labels[field_key] = value_label

        chart_wrap = tk.Frame(frame, bg=CARD)
        chart_wrap.pack(fill=tk.X, padx=12, pady=(0, 8))
        self.external_exposure_figure = Figure(figsize=(6.4, 2.4), dpi=100, facecolor=CARD)
        self.external_exposure_ax = self.external_exposure_figure.add_subplot(111)
        self.external_exposure_ax.set_facecolor(CARD)
        self.external_exposure_canvas = FigureCanvasTkAgg(self.external_exposure_figure, master=chart_wrap)
        self.external_exposure_canvas.get_tk_widget().pack(fill=tk.X, expand=True)

        table_wrap = ttk.Frame(frame, style="Dark.TFrame", padding=10)
        table_wrap.pack(fill=tk.BOTH, expand=True, padx=12, pady=(0, 10))
        table_wrap.grid_rowconfigure(0, weight=1)
        table_wrap.grid_columnconfigure(0, weight=1)

        self.external_exposure_tree = ttk.Treeview(
            table_wrap,
            columns=("Port", "Service", "Severity", "Reason"),
            show="headings",
            style="Dark.Treeview",
            selectmode="none",
            height=10,
        )
        self.external_exposure_tree.heading("Port", text="Port", anchor="center")
        self.external_exposure_tree.heading("Service", text="Service", anchor="w")
        self.external_exposure_tree.heading("Severity", text="Risk", anchor="center")
        self.external_exposure_tree.heading("Reason", text="Exposure Detail", anchor="w")
        self.external_exposure_tree.column("Port", width=90, anchor="center", stretch=False)
        self.external_exposure_tree.column("Service", width=140, anchor="w", stretch=False)
        self.external_exposure_tree.column("Severity", width=100, anchor="center", stretch=False)
        self.external_exposure_tree.column("Reason", width=620, anchor="w", stretch=True)
        self.external_exposure_tree.tag_configure("HIGH", background="#5C1F1F", foreground="#FFD7D7")
        self.external_exposure_tree.tag_configure("MEDIUM", background="#5A3A16", foreground="#FFE3BF")
        self.external_exposure_tree.tag_configure("LOW", background="#1B3D2F", foreground="#D6F5E3")
        self.external_exposure_tree.tag_configure("INFO", background="#1C3552", foreground="#D8ECFF")

        exposure_vsb = ttk.Scrollbar(table_wrap, orient="vertical", command=self.external_exposure_tree.yview)
        exposure_hsb = ttk.Scrollbar(table_wrap, orient="horizontal", command=self.external_exposure_tree.xview)
        self.external_exposure_tree.configure(yscrollcommand=exposure_vsb.set, xscrollcommand=exposure_hsb.set)

        self.external_exposure_tree.grid(row=0, column=0, sticky="nsew")
        exposure_vsb.grid(row=0, column=1, sticky="ns")
        exposure_hsb.grid(row=1, column=0, sticky="ew")

        self._render_external_exposure_view(None)

    def _render_external_exposure_view(self, scan: dict[str, Any] | None) -> None:
        if not hasattr(self, "external_exposure_tree") or self.external_exposure_tree is None:
            return

        for child in list(self.external_exposure_tree.get_children()):
            try:
                self.external_exposure_tree.delete(child)
            except Exception:
                pass

        empty_fields = {
            "ip_address": "-",
            "hostname": "-",
            "device_type": "-",
            "mac_vendor": "-",
            "risk_level": "-",
            "open_ports": "-",
            "scan_type": "-",
            "visibility": scanner.LIMITED_VISIBILITY_NOTE,
        }

        if not scan:
            self.external_exposure_title_label.config(text="Attacker View: No data loaded", fg=TEXT)
            self.external_exposure_note_label.config(
                text="Run Scan Without Credentials to populate attacker-view exposure data.",
                fg="#FFB74D",
            )
            for field_key, value in empty_fields.items():
                self.external_exposure_value_labels[field_key].config(text=value, fg=TEXT)
            self.external_exposure_tree.insert(
                "",
                tk.END,
                values=("-", "-", "INFO", "No unauthenticated exposure data is available for this device."),
                tags=("INFO",),
            )
            self.external_exposure_ax.clear()
            self.external_exposure_ax.set_facecolor(CARD)
            self.external_exposure_ax.text(0.5, 0.5, "No exposure data", ha="center", va="center", color=TEXT)
            self.external_exposure_figure.tight_layout(pad=1.0)
            self.external_exposure_canvas.draw_idle()
            return

        exposure = scan.get("external_exposure") or {}
        risk_level = str(exposure.get("risk_level") or scan.get("risk_level") or "UNKNOWN").upper()
        risk_color = {"HIGH": "#FF8A80", "MEDIUM": "#FFB74D", "LOW": "#81C784"}.get(risk_level, TEXT)
        open_ports = exposure.get("open_ports") or []
        mac_address = str(exposure.get("mac_address") or "").strip()
        mac_vendor = str(exposure.get("mac_vendor") or "").strip()
        mac_text = ", ".join(part for part in (mac_address, mac_vendor) if part) or "Unavailable"
        open_port_label = ", ".join(str(item.get("port")) for item in open_ports) or "None"

        self.external_exposure_title_label.config(
            text=f"Attacker View: {scanner.LIMITED_VISIBILITY_LABEL}",
            fg=TEXT,
        )
        self.external_exposure_note_label.config(
            text=str(exposure.get("limited_visibility_note") or scanner.LIMITED_VISIBILITY_NOTE),
            fg="#FFB74D",
        )

        field_values = {
            "ip_address": str(exposure.get("ip_address") or scan.get("target_ip") or "Unavailable"),
            "hostname": str(exposure.get("hostname") or scan.get("hostname") or "Unavailable"),
            "device_type": str(exposure.get("device_type_label") or exposure.get("device_type") or "Unknown"),
            "mac_vendor": mac_text,
            "risk_level": risk_level,
            "open_ports": f"{len(open_ports)} exposed ({open_port_label})" if open_ports else "0 exposed",
            "scan_type": str(exposure.get("scan_type") or scanner.LIMITED_VISIBILITY_LABEL),
            "visibility": str(exposure.get("limited_visibility_note") or scanner.LIMITED_VISIBILITY_NOTE),
        }
        for field_key, value in field_values.items():
            color = risk_color if field_key == "risk_level" else TEXT
            self.external_exposure_value_labels[field_key].config(text=value, fg=color)

        if open_ports:
            for item in open_ports:
                severity = str(item.get("severity") or "LOW").upper()
                self.external_exposure_tree.insert(
                    "",
                    tk.END,
                    values=(
                        str(item.get("port") or "-"),
                        str(item.get("service") or "Unknown"),
                        severity,
                        str(item.get("reason") or ""),
                    ),
                    tags=(severity,),
                )
        else:
            self.external_exposure_tree.insert(
                "",
                tk.END,
                values=("-", "-", "LOW", "No monitored services were exposed during the unauthenticated scan."),
                tags=("LOW",),
            )

        service_breakdown = exposure.get("service_breakdown") or {}
        self.external_exposure_ax.clear()
        self.external_exposure_ax.set_facecolor(CARD)
        if service_breakdown:
            labels = list(service_breakdown.keys())
            values = [int(service_breakdown[label]) for label in labels]
            bars = self.external_exposure_ax.bar(labels, values, color="#FF7043")
            self.external_exposure_ax.set_title("Exposed Services", color=TEXT, fontsize=10)
            self.external_exposure_ax.tick_params(axis="x", colors=TEXT, labelsize=8)
            self.external_exposure_ax.tick_params(axis="y", colors=TEXT, labelsize=8)
            for bar, value in zip(bars, values, strict=True):
                self.external_exposure_ax.text(
                    bar.get_x() + (bar.get_width() / 2),
                    bar.get_height() + 0.05,
                    str(value),
                    ha="center",
                    va="bottom",
                    color=TEXT,
                    fontsize=8,
                )
        else:
            self.external_exposure_ax.text(0.5, 0.5, "No exposed monitored services", ha="center", va="center", color=TEXT)
        self.external_exposure_figure.tight_layout(pad=1.0)
        self.external_exposure_canvas.draw_idle()

    def _make_software_inventory_tab(self, title: str) -> None:
        """
        Creates a scrollable inventory view for installed applications.
        """
        frame = tk.Frame(self.notebook, bg=CARD)
        self.notebook.add(frame, text=title)
        frame.pack_propagate(False)

        scroll_wrap = tk.Frame(frame, bg=CARD)
        scroll_wrap.pack(fill=tk.BOTH, expand=True)

        scroll_canvas = tk.Canvas(scroll_wrap, bg=CARD, highlightthickness=0, borderwidth=0)
        outer_vsb = ttk.Scrollbar(scroll_wrap, orient="vertical", command=scroll_canvas.yview)
        scroll_canvas.configure(yscrollcommand=outer_vsb.set)

        outer_vsb.pack(side=tk.RIGHT, fill=tk.Y, padx=(0, 6), pady=(10, 10))
        scroll_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 4), pady=(10, 10))

        content_frame = tk.Frame(scroll_canvas, bg=CARD)
        canvas_window = scroll_canvas.create_window((0, 0), window=content_frame, anchor="nw")

        def _sync_software_scroll(_event: Any = None) -> None:
            scroll_canvas.configure(scrollregion=scroll_canvas.bbox("all"))
            try:
                scroll_canvas.itemconfigure(canvas_window, width=scroll_canvas.winfo_width())
            except Exception:
                pass

        def _software_mousewheel(event: Any) -> None:
            delta = 0
            if getattr(event, "delta", 0):
                delta = -1 * int(event.delta / 120)
            elif getattr(event, "num", None) == 4:
                delta = -1
            elif getattr(event, "num", None) == 5:
                delta = 1
            if delta:
                scroll_canvas.yview_scroll(delta, "units")

        scroll_canvas.bind("<Configure>", _sync_software_scroll)
        content_frame.bind("<Configure>", _sync_software_scroll)
        scroll_canvas.bind("<MouseWheel>", _software_mousewheel)
        content_frame.bind("<MouseWheel>", _software_mousewheel)

        header = tk.Frame(content_frame, bg=CARD)
        header.pack(fill=tk.X, padx=12, pady=(10, 4))

        tk.Label(
            header,
            text="Third-Party Software Overview",
            bg=CARD,
            fg=TEXT,
            font=("Segoe UI", 12, "bold"),
            anchor="w",
            justify=tk.LEFT,
        ).pack(anchor="w", fill=tk.X)

        self.software_total_label = tk.Label(
            header,
            text="Total Applications: 0",
            bg=CARD,
            fg=MUTED,
            font=("Segoe UI", 10, "bold"),
            anchor="w",
            justify=tk.LEFT,
        )
        self.software_total_label.pack(anchor="w", fill=tk.X, pady=(6, 0))

        self.software_risk_summary_label = tk.Label(
            header,
            text="Risky Applications: 0",
            bg=CARD,
            fg="#FFB74D",
            font=("Segoe UI", 10, "bold"),
            anchor="w",
            justify=tk.LEFT,
        )
        self.software_risk_summary_label.pack(anchor="w", fill=tk.X, pady=(4, 0))

        self.software_status_label = tk.Label(
            header,
            text="Third-party Software: Secure (No risky applications found)",
            bg=CARD,
            fg=MUTED,
            font=("Segoe UI", 10),
            anchor="w",
            justify=tk.LEFT,
        )
        self.software_status_label.pack(anchor="w", fill=tk.X, pady=(4, 0))

        sections = tk.PanedWindow(content_frame, orient=tk.VERTICAL, bg=CARD, sashwidth=8, sashrelief=tk.RAISED, bd=0)
        sections.pack(fill=tk.BOTH, expand=True, padx=12, pady=(8, 10))

        risky_panel = tk.Frame(sections, bg=CARD)
        tk.Label(
            risky_panel,
            text="Risky Applications Detected",
            bg=CARD,
            fg=TEXT,
            font=("Segoe UI", 11, "bold"),
            anchor="w",
            justify=tk.LEFT,
        ).pack(anchor="w", fill=tk.X, pady=(0, 2))

        risky_container = ttk.Frame(risky_panel, style="Dark.TFrame", padding=10)
        risky_container.pack(fill=tk.BOTH, expand=True, pady=(0, 4))
        risky_container.grid_rowconfigure(0, weight=1)
        risky_container.grid_columnconfigure(0, weight=1)

        self.risky_software_tree = ttk.Treeview(
            risky_container,
            columns=("Application Name", "Version", "Risk", "Reason"),
            show="headings",
            style="Dark.Treeview",
            selectmode="none",
            height=4,
        )
        self.risky_software_tree.heading("Application Name", text="Application Name", anchor="w")
        self.risky_software_tree.heading("Version", text="Version", anchor="center")
        self.risky_software_tree.heading("Risk", text="Risk", anchor="center")
        self.risky_software_tree.heading("Reason", text="Reason", anchor="w")
        self.risky_software_tree.column("Application Name", width=260, anchor="w", stretch=True)
        self.risky_software_tree.column("Version", width=120, anchor="center", stretch=False)
        self.risky_software_tree.column("Risk", width=90, anchor="center", stretch=False)
        self.risky_software_tree.column("Reason", width=420, anchor="w", stretch=True)
        self.risky_software_tree.tag_configure("HIGH", background="#5C1F1F", foreground="#FFD7D7")
        self.risky_software_tree.tag_configure("MEDIUM", background="#5A3A16", foreground="#FFE3BF")
        self.risky_software_tree.tag_configure("LOW", background="#1B3D2F", foreground="#D6F5E3")

        risky_vsb = ttk.Scrollbar(risky_container, orient="vertical", command=self.risky_software_tree.yview)
        risky_hsb = ttk.Scrollbar(risky_container, orient="horizontal", command=self.risky_software_tree.xview)
        self.risky_software_tree.configure(yscrollcommand=risky_vsb.set, xscrollcommand=risky_hsb.set)

        self.risky_software_tree.grid(row=0, column=0, sticky="nsew")
        risky_vsb.grid(row=0, column=1, sticky="ns")
        risky_hsb.grid(row=1, column=0, sticky="ew")

        inventory_panel = tk.Frame(sections, bg=CARD)
        tk.Label(
            inventory_panel,
            text="Installed Applications (Full Inventory)",
            bg=CARD,
            fg=TEXT,
            font=("Segoe UI", 11, "bold"),
            anchor="w",
            justify=tk.LEFT,
        ).pack(anchor="w", fill=tk.X, pady=(0, 2))

        container = ttk.Frame(inventory_panel, style="Dark.TFrame", padding=10)
        container.pack(fill=tk.BOTH, expand=True, pady=(0, 4))
        container.grid_rowconfigure(0, weight=1)
        container.grid_columnconfigure(0, weight=1)

        self.software_tree = ttk.Treeview(
            container,
            columns=("Application Name", "Version"),
            show="headings",
            style="Dark.Treeview",
            selectmode="none",
            height=15,
        )
        self.software_tree.heading("Application Name", text="Application Name", anchor="w")
        self.software_tree.heading("Version", text="Version", anchor="center")
        self.software_tree.column("Application Name", width=300, anchor="w", stretch=True)
        self.software_tree.column("Version", width=150, anchor="center", stretch=False)
        self.software_tree.tag_configure("oddrow", background="#181818")
        self.software_tree.tag_configure("evenrow", background="#202020")

        vsb = ttk.Scrollbar(container, orient="vertical", command=self.software_tree.yview)
        hsb = ttk.Scrollbar(container, orient="horizontal", command=self.software_tree.xview)
        self.software_tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.software_tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

        sections.add(risky_panel, minsize=110, stretch="never")
        sections.add(inventory_panel, minsize=220, stretch="always")

    def _make_live_applications_tab(self, title: str) -> None:
        frame = tk.Frame(self.notebook, bg=CARD)
        self.notebook.add(frame, text=title)

        header = tk.Frame(frame, bg=CARD)
        header.pack(fill=tk.X, padx=12, pady=(10, 6))

        tk.Label(
            header,
            text="Live Application Monitor",
            bg=CARD,
            fg=TEXT,
            font=("Segoe UI", 12, "bold"),
        ).pack(side=tk.LEFT, anchor="w")

        self.live_apps_refresh_btn = tk.Button(
            header,
            text="Refresh",
            command=self._refresh_live_applications_clicked,
            bg="#2C2C2C",
            fg=TEXT,
            relief=tk.FLAT,
        )
        self.live_apps_refresh_btn.pack(side=tk.RIGHT)

        self.live_apps_status_label = tk.Label(
            frame,
            text="Select a remote-scanned device to monitor running processes.",
            bg=CARD,
            fg=MUTED,
            anchor="w",
            justify=tk.LEFT,
            wraplength=960,
        )
        self.live_apps_status_label.pack(fill=tk.X, padx=12, pady=(0, 8))

        container = ttk.Frame(frame, style="Dark.TFrame", padding=10)
        container.pack(fill=tk.BOTH, expand=True, padx=12, pady=(0, 10))
        container.grid_rowconfigure(0, weight=1)
        container.grid_columnconfigure(0, weight=1)

        self.live_apps_tree = ttk.Treeview(
            container,
            columns=("Process Name", "PID", "Memory Usage (MB)", "CPU Usage"),
            show="headings",
            style="Dark.Treeview",
            selectmode="none",
            height=18,
        )
        self.live_apps_tree.heading("Process Name", text="Process Name", anchor="w")
        self.live_apps_tree.heading("PID", text="PID", anchor="center")
        self.live_apps_tree.heading("Memory Usage (MB)", text="Memory Usage (MB)", anchor="center")
        self.live_apps_tree.heading("CPU Usage", text="CPU Usage", anchor="center")
        self.live_apps_tree.column("Process Name", width=320, anchor="w", stretch=True)
        self.live_apps_tree.column("PID", width=100, anchor="center", stretch=False)
        self.live_apps_tree.column("Memory Usage (MB)", width=160, anchor="center", stretch=False)
        self.live_apps_tree.column("CPU Usage", width=140, anchor="center", stretch=False)
        self.live_apps_tree.tag_configure("oddrow", background="#181818")
        self.live_apps_tree.tag_configure("evenrow", background="#202020")

        live_vsb = ttk.Scrollbar(container, orient="vertical", command=self.live_apps_tree.yview)
        live_hsb = ttk.Scrollbar(container, orient="horizontal", command=self.live_apps_tree.xview)
        self.live_apps_tree.configure(yscrollcommand=live_vsb.set, xscrollcommand=live_hsb.set)

        self.live_apps_tree.grid(row=0, column=0, sticky="nsew")
        live_vsb.grid(row=0, column=1, sticky="ns")
        live_hsb.grid(row=1, column=0, sticky="ew")

        self._populate_live_processes_table([], "No live application data loaded yet.")
        self.live_apps_refresh_btn.config(state=tk.DISABLED)

    def _current_detail_tab_title(self) -> str:
        try:
            return str(self.notebook.tab(self.notebook.select(), "text") or "")
        except Exception:
            return ""

    def _set_live_applications_status(self, message: str, color: str = MUTED) -> None:
        if hasattr(self, "live_apps_status_label") and self.live_apps_status_label is not None:
            self.live_apps_status_label.config(text=message, fg=color)

    def _populate_live_processes_table(self, processes: list[dict[str, Any]], empty_message: str) -> None:
        if not hasattr(self, "live_apps_tree") or self.live_apps_tree is None:
            return

        for child in list(self.live_apps_tree.get_children()):
            try:
                self.live_apps_tree.delete(child)
            except Exception:
                pass

        if not processes:
            self.live_apps_tree.insert("", tk.END, values=(empty_message, "-", "-", "-"), tags=("evenrow",))
            return

        for index, item in enumerate(processes):
            memory_mb = item.get("memory_mb")
            cpu_usage = item.get("cpu")
            row_tag = "evenrow" if index % 2 == 0 else "oddrow"
            self.live_apps_tree.insert(
                "",
                tk.END,
                values=(
                    str(item.get("name") or ""),
                    str(item.get("pid") or "-"),
                    f"{float(memory_mb):.1f}" if memory_mb is not None else "-",
                    f"{float(cpu_usage):.2f}" if cpu_usage is not None else "-",
                ),
                tags=(row_tag,),
            )

    def _render_live_applications_view(self) -> None:
        if not hasattr(self, "live_apps_tree") or self.live_apps_tree is None:
            return

        device = self.selected_device
        if not device:
            self.live_apps_refresh_btn.config(state=tk.DISABLED)
            self._populate_live_processes_table([], "Select a device to view live applications.")
            self._set_live_applications_status("Select a device to view live applications.")
            return

        if self.selected_scan and risk.is_limited_visibility_scan(self.selected_scan):
            self.live_apps_refresh_btn.config(state=tk.DISABLED)
            self._populate_live_processes_table([], "Live applications are not available in unauthenticated mode.")
            self._set_live_applications_status(
                scanner.LIMITED_VISIBILITY_NOTE,
                color="#FFB74D",
            )
            return

        if self._device_type(device) == "Mobile":
            self.live_apps_refresh_btn.config(state=tk.DISABLED)
            self._populate_live_processes_table([], "Mobile devices are not supported for live application monitoring.")
            self._set_live_applications_status(
                "Mobile devices are not supported for scanning or live application monitoring.",
                color="#FFB74D",
            )
            return

        ip_address = str(device.get("ip") or "").strip()
        display_name = self._device_display_name(device)
        cache_key = ip_address or self._device_key(device)
        processes = self.live_process_cache.get(cache_key, [])
        last_refreshed = self.live_process_last_refreshed.get(cache_key)
        is_local_scan = bool(
            isinstance(self.selected_scan, dict)
            and str(self.selected_scan.get("connection_status") or "").strip().lower() == "local scan"
        )
        is_connected = bool(ip_address and self._get_connected_device_entry(ip_address))

        can_refresh = bool((ip_address and is_connected) or is_local_scan) and not self.live_process_fetch_in_progress
        self.live_apps_refresh_btn.config(state=tk.NORMAL if can_refresh else tk.DISABLED)

        if self.live_process_fetch_in_progress:
            self._set_live_applications_status(f"Fetching running applications for {display_name}...", color=MUTED)
            self._populate_live_processes_table([], "Loading running applications...")
            return

        if not ip_address:
            self._populate_live_processes_table([], "No IP address available for this device.")
            self._set_live_applications_status("No IP address available for this device.", color="#FFB74D")
            return

        if last_refreshed is not None:
            if processes:
                self._populate_live_processes_table(processes, "No active processes found.")
                self._set_live_applications_status(
                    f"Showing {len(processes)} running processes for {display_name}. Last refreshed: {last_refreshed}.",
                    color=MUTED,
                )
            else:
                self._populate_live_processes_table([], "No active processes found.")
                self._set_live_applications_status(
                    f"No active processes found for {display_name}. Last refreshed: {last_refreshed}.",
                    color=MUTED,
                )
            return

        self._populate_live_processes_table([], "No live application data loaded yet.")
        if is_local_scan:
            self._set_live_applications_status(
                f"Open this tab or click Refresh to fetch running applications for {display_name}.",
                color=MUTED,
            )
            return
        if self._is_local_device(device):
            self._set_live_applications_status(
                f"Run Local Scan for {display_name} to enable live application monitoring.",
                color="#FFB74D",
            )
            return
        if is_connected:
            self._set_live_applications_status(
                f"Open this tab or click Refresh to fetch running applications for {display_name}.",
                color=MUTED,
            )
        else:
            self._set_live_applications_status(
                f"Use Connect & Scan once for {display_name} to enable live application monitoring.",
                color="#FFB74D",
            )

    def _on_detail_tab_changed(self, _event: Any) -> None:
        if self._current_detail_tab_title() != "Live Applications":
            return
        self._render_live_applications_view()
        self._request_live_application_refresh(auto_trigger=True)

    def _refresh_live_applications_clicked(self) -> None:
        self._request_live_application_refresh(auto_trigger=False)

    def _request_live_application_refresh(self, auto_trigger: bool = False) -> None:
        if self.live_process_fetch_in_progress:
            return

        device = self.selected_device
        if not device:
            self._populate_live_processes_table([], "Select a device to view live applications.")
            self._set_live_applications_status("Select a device to view live applications.", color="#FFB74D")
            return

        if self.selected_scan and risk.is_limited_visibility_scan(self.selected_scan):
            message = scanner.LIMITED_VISIBILITY_NOTE
            self.live_apps_refresh_btn.config(state=tk.DISABLED)
            self._populate_live_processes_table([], "Live applications are not available in unauthenticated mode.")
            self._set_live_applications_status(message, color="#FFB74D")
            if not auto_trigger:
                messagebox.showinfo("Live Applications", message)
            return

        if self._device_type(device) == "Mobile":
            message = "Mobile devices are not supported for scanning or live application monitoring."
            self.live_apps_refresh_btn.config(state=tk.DISABLED)
            self._populate_live_processes_table([], "Mobile devices are not supported for live application monitoring.")
            self._set_live_applications_status(message, color="#FFB74D")
            if not auto_trigger:
                messagebox.showinfo("Live Applications", message)
            return

        ip_address = str(device.get("ip") or "").strip()
        display_name = self._device_display_name(device)
        is_local_scan = bool(
            isinstance(self.selected_scan, dict)
            and str(self.selected_scan.get("connection_status") or "").strip().lower() == "local scan"
        )
        if not ip_address:
            if not is_local_scan:
                self._populate_live_processes_table([], "No IP address available for this device.")
                self._set_live_applications_status("No IP address available for this device.", color="#FFB74D")
                return

        if self._is_local_device(device) and not is_local_scan:
            self._populate_live_processes_table([], "Run a local scan to load live applications.")
            message = f"Run Local Scan once for {display_name} to enable live application monitoring."
            self._set_live_applications_status(message, color="#FFB74D")
            if not auto_trigger:
                messagebox.showinfo("Live Applications", message)
            return

        connection_entry = self._get_connected_device_entry(ip_address) if not is_local_scan else None
        if not is_local_scan and not connection_entry:
            self._populate_live_processes_table([], "Live monitoring credentials are not available.")
            message = f"Use Connect & Scan once for {display_name} to enable live application monitoring."
            self._set_live_applications_status(message, color="#FFB74D")
            if not auto_trigger:
                messagebox.showinfo("Live Applications", message)
            return

        self.live_process_fetch_in_progress = True
        self.live_apps_refresh_btn.config(state=tk.DISABLED)
        self._populate_live_processes_table([], "Loading running applications...")
        self._set_live_applications_status(f"Fetching running applications for {display_name}...", color=MUTED)
        threading.Thread(
            target=self._live_process_worker,
            args=(ip_address, connection_entry["session"] if connection_entry else None, is_local_scan),
            daemon=True,
        ).start()

    def _live_process_worker(self, ip_address: str, session: Any | None, is_local_scan: bool) -> None:
        try:
            if is_local_scan:
                processes = scanner.fetch_local_processes(limit=50)
            else:
                processes = scanner.fetch_live_processes_with_session(ip_address, session, limit=50)
            fetched_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            self.after(
                0,
                lambda ip_address=ip_address, processes=processes, fetched_at=fetched_at: self._on_live_process_success(
                    ip_address,
                    processes,
                    fetched_at,
                ),
            )
        except scanner.RemoteScanError as exc:
            self.after(
                0,
                lambda ip_address=ip_address, message=str(exc), code=exc.code: self._on_live_process_failure(
                    ip_address,
                    scanner.RemoteScanError(message, code=code),
                ),
            )
        except Exception as exc:
            self.after(
                0,
                lambda ip_address=ip_address, message=str(exc): self._on_live_process_failure(
                    ip_address,
                    scanner.RemoteScanError(message, code="execution_error"),
                ),
            )

    def _on_live_process_success(self, ip_address: str, processes: list[dict[str, Any]], fetched_at: str) -> None:
        self.live_process_fetch_in_progress = False
        self.live_process_cache[ip_address] = processes
        self.live_process_last_refreshed[ip_address] = fetched_at
        if self.selected_device and str(self.selected_device.get("ip") or "").strip() == ip_address:
            self._render_live_applications_view()
        else:
            self._render_live_applications_view()

    def _on_live_process_failure(self, ip_address: str, error: scanner.RemoteScanError) -> None:
        self.live_process_fetch_in_progress = False
        if error.code in {"invalid_credentials", "connection_failed", "timeout"}:
            self.connected_devices.pop(ip_address, None)

        if self.selected_device and str(self.selected_device.get("ip") or "").strip() == ip_address:
            self._update_connection_status_ui()
            self._populate_live_processes_table([], "Unable to fetch running applications.")
            if error.code in {"invalid_credentials", "connection_failed", "timeout"}:
                self._set_live_applications_status(
                    "Session expired, please reconnect before refreshing live applications.",
                    color="#FF8A80",
                )
            else:
                self._set_live_applications_status("Unable to fetch running applications.", color="#FF8A80")
        else:
            self._update_connection_status_ui()

    # Threat Analysis tab removed: Threat Analysis now renders only in the main dashboard area.

    def _refresh_all(self) -> None:
        selected_key = self._device_key(self.selected_device) if self.selected_device else ""
        self._clear_saved_scan_snapshots()
        self._sync_devices_from_results()
        self._populate_system_list()
        self._update_summary_label()
        if self.device_rows:
            target_key = selected_key or self._device_key(self.device_rows[0])
            self._select_device_by_key(target_key)
        else:
            self.selected_device = None
            self.selected_scan = None
            self.selected_unauth_scan = None
            self._update_connection_status_ui()

    def _clear_saved_devices(self) -> None:
        if self._any_scan_in_progress():
            messagebox.showinfo("Clear Saved Devices", "Wait for the current scan to finish before clearing saved devices.")
            return

        confirmed = messagebox.askyesno(
            "Clear Saved Devices",
            "Remove all saved scan history and remembered remote devices from the dashboard? Only this device will remain until you scan again.",
        )
        if not confirmed:
            return

        scanner.save_scan_results([])
        scanner.save_unauthenticated_scan_results([])
        self.results = []
        self.unauth_results = []
        self.devices = [self._make_local_device_placeholder()]
        self.device_rows = []
        self.selected_device = None
        self.selected_scan = None
        self.selected_unauth_scan = None
        self.connected_devices.clear()
        self.live_process_cache.clear()
        self.live_process_last_refreshed.clear()
        self._reset_remote_scan_progress_ui()
        self._populate_system_list()
        self._update_summary_label()
        self._populate_live_processes_table([], "Run a local scan to load live applications.")
        self._set_live_applications_status("Run a local scan to load live applications.", color=MUTED)
        self._render_external_exposure_view(None)
        self.discovery_status_label.config(
            text="Saved devices cleared. Showing only this device until you scan or discover devices again."
        )
        if self.device_rows:
            self._select_device_by_key(self._device_key(self.device_rows[0]))
        else:
            self._update_connection_status_ui()

    def _populate_system_list(self) -> None:
        self.system_list.delete(0, tk.END)
        self.device_rows = list(self.devices)
        for device in self.device_rows:
            display_name = self._device_display_name(device)
            device_type = self._device_type(device)
            status = str(device.get("status") or ("Online" if device.get("online") else "Offline"))
            auth_scan = device.get("scan_result") or {}
            unauth_scan = device.get("unauth_scan_result") or {}
            if auth_scan:
                scan = auth_scan
                lvl = str(scan.get("risk_level") or "UNKNOWN").upper()
                score = int(scan.get("risk_score") or 0)
                line = f"{display_name} ({device_type}) | {status:<18} | {lvl:<8} | {score:>3}"
            elif unauth_scan:
                lvl = str(unauth_scan.get("risk_level") or "UNKNOWN").upper()
                score = int(unauth_scan.get("risk_score") or 0)
                line = f"{display_name} ({device_type}) | {'Exposure Assessed':<18} | {'EXT ' + lvl:<11} | {score:>3}"
            else:
                line = f"{display_name} ({device_type}) | {status:<18} | {'NOT SCANNED':<11} | {'-':>3}"
            self.system_list.insert(tk.END, line)
            try:
                self.system_list.itemconfig(self.system_list.size() - 1, fg=self._device_type_color(device_type))
            except Exception:
                pass

    def _update_component_charts(self, scan: dict[str, Any]) -> None:
        scores = risk.compute_component_scores(scan)
        contrib = risk.compute_contributions(scores)
        total_score = sum(scores.values())
        self.component_score_cache = scores
        self.component_contrib_cache = contrib
        self.pie_click_info.config(text="" if total_score > 0 else "System is secure. No risks detected.")

        labels = []
        sizes = []
        colors = []
        for comp in risk.COMPONENTS:
            value = float(contrib.get(comp, 0.0))
            if value > 0:
                labels.append(comp)
                sizes.append(value)
                colors.append(risk.COMPONENT_COLORS.get(comp, ACCENT))

        self.ax_pie.clear()
        self.ax_bar.clear()
        self.ax_pie.set_facecolor(CARD)
        self.ax_bar.set_facecolor(CARD)

        if sizes:
            wedges, _texts, _auto = self.ax_pie.pie(
                sizes,
                labels=labels,
                colors=colors,
                autopct="%1.1f%%",
                startangle=90,
                textprops={"color": TEXT, "fontsize": 8},
            )
            self.ax_pie.set_title("Contribution Pie", color=TEXT, fontsize=10)
            self.ax_pie.axis("equal")
            self._pie_wedges = list(zip(wedges, labels, sizes, strict=True))
        else:
            self.ax_pie.text(0.5, 0.5, "System is secure.\nNo risks detected.", ha="center", va="center", color=TEXT)
            self._pie_wedges = []

        # Bar chart with highest contributor highlighted
        ordered = sorted([(k, float(contrib.get(k, 0.0))) for k in risk.COMPONENTS], key=lambda x: x[1], reverse=True)
        bar_labels = [x[0] for x in ordered]
        bar_values = [x[1] for x in ordered]
        if bar_values:
            max_val = max(bar_values)
            bar_colors = []
            for name, val in ordered:
                base = risk.COMPONENT_COLORS.get(name, "#4FC3F7")
                if val == max_val and max_val > 0:
                    bar_colors.append("#FF5252")  # highlight top contributor
                else:
                    bar_colors.append(base)
            bars = self.ax_bar.barh(bar_labels, bar_values, color=bar_colors)
            self.ax_bar.invert_yaxis()
            self.ax_bar.set_xlim(0, 100)
            self.ax_bar.set_xlabel("% Contribution", color=TEXT, fontsize=9)
            self.ax_bar.set_title("Contribution Bar", color=TEXT, fontsize=10)
            self.ax_bar.tick_params(axis="x", colors=TEXT, labelsize=8)
            self.ax_bar.tick_params(axis="y", colors=TEXT, labelsize=8)
            for b in bars:
                w = b.get_width()
                self.ax_bar.text(w + 1, b.get_y() + b.get_height() / 2, f"{w:.1f}%", va="center", color=TEXT, fontsize=8)

        for spine in self.ax_bar.spines.values():
            spine.set_color("#444444")

        self.figure.tight_layout(pad=1.2)
        self.canvas.draw_idle()

    def _update_summary_label(self) -> None:
        counts, _ = risk.risk_distribution(self.results)
        total = len(self.results)
        unauth_total = len(self.unauth_results)
        txt = (
            f"Total: {total}  |  CRITICAL: {counts['CRITICAL']}  HIGH: {counts['HIGH']}  "
            f"MEDIUM: {counts['MEDIUM']}  LOW: {counts['LOW']}"
        )
        if unauth_total:
            txt += f"  |  External Exposure Scans: {unauth_total}"
        if counts["CRITICAL"] or counts["HIGH"]:
            txt += "  -> Immediate action required for high-risk systems."
        self.summary_label.config(text=txt)

    def _render_selected_scan(self) -> None:
        if not self.selected_scan:
            return
        scan = self.selected_scan
        exposure_scan = self.selected_unauth_scan
        if exposure_scan is None and risk.is_limited_visibility_scan(scan):
            exposure_scan = scan
        scores = risk.compute_component_scores(scan)
        contrib = risk.compute_contributions(scores)
        self._update_component_charts(scan)

        highest = max(scores.items(), key=lambda kv: kv[1])[0] if scores else ""
        for comp in risk.COMPONENTS:
            s = int(scores.get(comp, 0))
            c = float(contrib.get(comp, 0.0))
            self.component_bars[comp]["value"] = s
            self.component_labels[comp].config(text=f"{s} pts ({c:.0f}%)")
            if comp == highest and s > 0:
                self.component_labels[comp].config(fg="#FF8A80")
            else:
                self.component_labels[comp].config(fg=MUTED)

        self.why_label.config(text=risk.summarize_risk_reason(scan))
        drivers = risk.top_risk_drivers(scan, limit=3)
        if drivers and drivers[0][1] > 0:
            lines = ["Top Risk Drivers:"]
            visible_drivers = [(name, points, pct) for name, points, pct in drivers if points > 0]
            for idx, (name, points, pct) in enumerate(visible_drivers, start=1):
                lines.append(f"{idx}. {name} -> {points} pts ({pct:.1f}%)")
            self.top_drivers_label.config(text="\n".join(lines), fg=TEXT)
        else:
            self.top_drivers_label.config(text="Top Risk Drivers: System is secure. No risks detected.", fg=MUTED)

        self._set_text(self.tab_antivirus, self._build_antivirus_text(scan))
        self._set_text(self.tab_firewall, self._build_firewall_text(scan))
        self._set_text(self.tab_ports, self._build_ports_text(scan))
        self._render_external_exposure_view(exposure_scan)
        self._set_text(self.tab_patching, self._build_patching_text(scan))
        self._render_installed_apps_table(scan)
        self._render_live_applications_view()
        self._set_text(self.tab_raw, json.dumps(scan, indent=2, ensure_ascii=False))
        if getattr(self, "right_view_mode", "risk") in ("threat", "split"):
            self._render_threat_analysis(scan)
        if self._current_detail_tab_title() == "Live Applications":
            self._request_live_application_refresh(auto_trigger=True)

    def _set_text(self, widget: tk.Text, value: str) -> None:
        widget.config(state=tk.NORMAL)
        widget.delete("1.0", tk.END)
        widget.insert(tk.END, value)
        widget.config(state=tk.DISABLED)

    def _build_antivirus_text(self, scan: dict[str, Any]) -> str:
        if risk.is_limited_visibility_scan(scan):
            return (
                f"{scanner.LIMITED_VISIBILITY_LABEL}\n"
                f"{scanner.LIMITED_VISIBILITY_NOTE}\n\n"
                "Antivirus status is not visible in attacker-view mode."
            )
        av = scan.get("antivirus") or {}
        av_score = int(risk.compute_component_scores(scan).get("Antivirus", 0))
        secure_status = "Antivirus: Secure (No risk detected)" if av_score == 0 else "Antivirus: Risk detected"
        lines = [
            secure_status,
            f"Status: {av.get('status', 'UNKNOWN')}",
            f"Real-time protection: {'ENABLED' if av.get('realtime_protection') is True else ('DISABLED' if av.get('realtime_protection') is False else 'UNKNOWN')}",
            f"Products: {', '.join(av.get('products') or []) or 'None detected'}",
            "",
            "Risk explanation:",
            "If antivirus is disabled, missing, or real-time protection is off, malware can run with less resistance.",
        ]
        return "\n".join(lines)

    def _build_firewall_text(self, scan: dict[str, Any]) -> str:
        if risk.is_limited_visibility_scan(scan):
            return (
                f"{scanner.LIMITED_VISIBILITY_LABEL}\n"
                f"{scanner.LIMITED_VISIBILITY_NOTE}\n\n"
                "Firewall profile state is not visible without authentication."
            )
        fw = scan.get("firewall") or {}
        profiles = fw.get("profiles") or {}
        lines = ["Profile status:"]
        for p in ("Domain", "Private", "Public"):
            st = (profiles.get(p) or {}).get("state", "UNKNOWN")
            marker = "OFF <- needs attention" if str(st).upper() == "OFF" else st
            lines.append(f"- {p}: {marker}")
        lines.extend(
            [
                "",
                "Why it matters:",
                "Firewall blocks unwanted inbound traffic. OFF profiles increase attack surface.",
            ]
        )
        return "\n".join(lines)

    def _build_ports_text(self, scan: dict[str, Any]) -> str:
        ports = scan.get("open_ports") or {}
        listening = ports.get("listening_ports") or []
        risky = ports.get("risky_open") or []
        if risk.is_limited_visibility_scan(scan):
            lines = [
                f"Scan Type: {scanner.LIMITED_VISIBILITY_LABEL}",
                scanner.LIMITED_VISIBILITY_NOTE,
                "",
                f"Monitored ports scanned: {', '.join(str(port) for port in sorted(scanner.UNAUTHENTICATED_PORT_MAP))}",
                f"Exposed monitored ports: {len(listening)}",
                f"Open ports: {', '.join(map(str, listening)) or 'None'}",
                "",
                "Exposure findings:",
            ]
            if risky:
                for item in risky:
                    lines.append(
                        f"- Port {item.get('port')} ({item.get('service') or 'Service'} | {item.get('severity')}): {item.get('reason')}"
                    )
            else:
                lines.append("- No monitored ports were exposed.")
            return "\n".join(lines)
        lines = [
            f"Total listening ports: {len(listening)}",
            f"Listening ports: {', '.join(map(str, listening[:50])) or 'None'}",
            "",
            "Risky open ports:",
        ]
        if risky:
            for item in risky:
                lines.append(f"- Port {item.get('port')} ({item.get('severity')}): {item.get('reason')}")
        else:
            lines.append("- None from monitored list (3389, 445, 21, 22)")
        return "\n".join(lines)

    def _build_patching_text(self, scan: dict[str, Any]) -> str:
        if risk.is_limited_visibility_scan(scan):
            return (
                f"{scanner.LIMITED_VISIBILITY_LABEL}\n"
                f"{scanner.LIMITED_VISIBILITY_NOTE}\n\n"
                "Patch inventory is not visible without authentication."
            )
        patches = scan.get("patches") or {}
        dt = patches.get("most_recent_patch_date")
        date_str = dt if isinstance(dt, str) else (dt.strftime("%Y-%m-%d") if dt else "Unknown")
        lines = [
            f"Last update date: {date_str}",
            f"Days since update: {patches.get('days_since_last_patch', 'Unknown')}",
            f"Total patches: {patches.get('total_patches', 0)}",
            f"Warning: {patches.get('patch_warning') or 'None'}",
            "",
            "Risk if outdated:",
            "Systems missing recent patches can be exploited using known vulnerabilities.",
        ]
        return "\n".join(lines)

    def _render_installed_apps_table(self, scan: dict[str, Any]) -> None:
        """
        Renders a scrollable full installed-application inventory (from JSON if present).
        """
        if risk.is_limited_visibility_scan(scan):
            if hasattr(self, "software_tree") and self.software_tree is not None:
                for child in list(self.software_tree.get_children()):
                    try:
                        self.software_tree.delete(child)
                    except Exception:
                        pass
                self.software_tree.insert(
                    "",
                    tk.END,
                    values=("Authentication Required", scanner.LIMITED_VISIBILITY_LABEL),
                    tags=("evenrow",),
                )
            if hasattr(self, "risky_software_tree") and self.risky_software_tree is not None:
                for child in list(self.risky_software_tree.get_children()):
                    try:
                        self.risky_software_tree.delete(child)
                    except Exception:
                        pass
                self.risky_software_tree.insert(
                    "",
                    tk.END,
                    values=(
                        "Limited Data",
                        "-",
                        "LOW",
                        scanner.LIMITED_VISIBILITY_NOTE,
                    ),
                    tags=("LOW",),
                )
            if hasattr(self, "software_total_label"):
                self.software_total_label.config(text="Total Applications: Limited Visibility")
            if hasattr(self, "software_risk_summary_label"):
                self.software_risk_summary_label.config(text="Risky Applications: Authentication Required")
            if hasattr(self, "software_status_label"):
                self.software_status_label.config(
                    text=f"Third-party Software: {scanner.LIMITED_VISIBILITY_NOTE}",
                    fg="#FFB74D",
                )
            return

        apps = scan.get("installed_apps") or []
        had_installed_apps_field = "installed_apps" in scan and bool(scan.get("installed_apps"))

        # Backward compatibility: older scan_results.json won't have installed_apps.
        if not apps:
            try:
                from . import installed_apps as installed_apps_module

                apps = installed_apps_module.get_installed_apps()
            except Exception:
                apps = []
            # Persist so JSON output / exports reflect the full inventory.
            try:
                scan["installed_apps"] = apps
                if not had_installed_apps_field:
                    scanner.save_scan_results(self.results)
            except Exception:
                pass

        analysis = risk.analyze_software_risk(apps)
        risky_apps = risk.filter_risky_apps(analysis)
        third_party = scan.setdefault("third_party_software", {})
        needs_save = False
        if scan.get("risky_apps") != risky_apps:
            scan["risky_apps"] = risky_apps
            needs_save = True
        if third_party.get("risky_apps") != risky_apps:
            third_party["risky_apps"] = risky_apps
            needs_save = True
        if needs_save:
            try:
                scanner.save_scan_results(self.results)
            except Exception:
                pass

        if not hasattr(self, "software_tree") or self.software_tree is None:
            return

        # Clear old rows.
        for child in list(self.software_tree.get_children()):
            try:
                self.software_tree.delete(child)
            except Exception:
                pass
        for child in list(self.risky_software_tree.get_children()):
            try:
                self.risky_software_tree.delete(child)
            except Exception:
                pass

        self.software_total_label.config(text=f"Total Applications: {len(apps)}")
        self.software_risk_summary_label.config(text=f"Risky Applications: {len(risky_apps)}")
        if risky_apps:
            self.software_status_label.config(text="Third-party Software: Risk detected", fg="#FFB74D")
        else:
            self.software_status_label.config(text="Third-party Software: Secure (No risky applications found)", fg=MUTED)

        severity_order = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
        if risky_apps:
            self.risky_software_tree["height"] = min(6, max(3, len(risky_apps)))
            for item in sorted(
                risky_apps,
                key=lambda row: (
                    severity_order.get(str(row.get("risk_level") or "").upper(), 9),
                    str(row.get("name") or "").lower(),
                ),
            ):
                risk_level = str(item.get("risk_level") or "LOW").upper()
                self.risky_software_tree.insert(
                    "",
                    tk.END,
                    values=(
                        str(item.get("name") or ""),
                        str(item.get("version") or "N/A"),
                        risk_level,
                        str(item.get("reason") or ""),
                    ),
                    tags=(risk_level,),
                )
        else:
            self.risky_software_tree["height"] = 1
            self.risky_software_tree.insert(
                "",
                tk.END,
                values=(
                    "No risky applications detected",
                    "-",
                    "LOW",
                    "Current heuristic checks did not flag any installed software.",
                ),
                tags=("LOW",),
            )

        for index, app in enumerate(apps):
            name = str(app.get("name") or "").strip()
            if not name:
                continue
            version = str(app.get("version") or "N/A").strip() or "N/A"
            row_tag = "evenrow" if index % 2 == 0 else "oddrow"
            self.software_tree.insert("", tk.END, values=(name, version), tags=(row_tag,))

    def _build_software_text(self, scan: dict[str, Any]) -> str:
        """
        Text fallback for environments where the table view isn't available.
        """
        if risk.is_limited_visibility_scan(scan):
            return (
                f"{scanner.LIMITED_VISIBILITY_LABEL}\n"
                f"{scanner.LIMITED_VISIBILITY_NOTE}\n\n"
                "Installed application inventory is not available in attacker-view mode."
            )
        apps = scan.get("installed_apps") or []
        risky_apps = risk.filter_risky_apps(risk.analyze_software_risk(apps))
        if not apps:
            return "Total Applications: 0"
        lines = [f"Total Applications: {len(apps)}", f"Risky Applications: {len(risky_apps)}", ""]
        if risky_apps:
            lines.append("Risky Applications Detected:")
            for item in risky_apps:
                lines.append(
                    f"- {item.get('name')} | {item.get('version')} | {item.get('risk_level')} | {item.get('reason')}"
                )
            lines.append("")
        for app in apps:
            name = str(app.get("name") or "").strip()
            if not name:
                continue
            version = str(app.get("version") or "N/A").strip() or "N/A"
            lines.append(f"{name}\t{version}")
        return "\n".join(lines)

    def _run_scan_clicked(self) -> None:
        try:
            new_scan = scanner.add_scan_to_results()
            self.results = scanner.load_scan_results()
            self._refresh_all()
            self._reset_remote_scan_progress_ui()
            self.discovery_status_label.config(text="Local scan completed.")
            self._select_device_by_key(scanner.get_device_id(new_scan))
        except Exception as e:
            messagebox.showerror("Error", f"Failed to run scan: {e}")

    def _reload_results(self) -> None:
        self.results = scanner.load_scan_results()
        self.unauth_results = scanner.load_unauthenticated_scan_results()
        self._refresh_all()
        self._reset_remote_scan_progress_ui()
        self.discovery_status_label.config(text="Reloaded saved scan results.")

    def _generate_word_report(self) -> None:
        try:
            path = build_word_report(self.results, output_path=Path("report.docx"))
            messagebox.showinfo("Report Generated", f"Word report written to: {path}")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to generate Word report: {e}")

    def _on_system_select(self, _event: Any) -> None:
        idxs = list(self.system_list.curselection())
        if not idxs:
            return
        self._apply_device_selection(idxs[0])

    def _on_pie_click(self, event: Any) -> None:
        if not hasattr(self, "_pie_wedges") or event.inaxes != self.ax_pie:
            return
        for wedge, label, _size in self._pie_wedges:
            if wedge.contains_point((event.x, event.y)):
                points = int(self.component_score_cache.get(label, 0))
                pct = float(self.component_contrib_cache.get(label, 0.0))
                self.pie_click_info.config(
                    text=f"{label}: {pct:.1f}% contribution | {points} points | Shows how much this component drives total risk."
                )
                break

    def _update_threat_charts(self, threats: list[dict[str, Any]]) -> None:
        """
        Threat-only charts (count-based):
        - Severity distribution: HIGH/MEDIUM/LOW
        - Category breakdown: Firewall/Open Ports/Services/Patching/Antivirus
        """
        if not hasattr(self, "threat_ax_sev") or not hasattr(self, "threat_ax_cat"):
            return

        # Clear prior charts
        self.threat_ax_sev.clear()
        self.threat_ax_cat.clear()
        self.threat_ax_sev.set_facecolor(CARD)
        self.threat_ax_cat.set_facecolor(CARD)

        if not threats:
            self.threat_ax_sev.text(0.5, 0.5, "No threat data", ha="center", va="center", color=TEXT, transform=self.threat_ax_sev.transAxes)
            self.threat_ax_cat.text(0.5, 0.5, "No threat data", ha="center", va="center", color=TEXT, transform=self.threat_ax_cat.transAxes)
            self.threat_canvas.draw_idle()
            return

        color_map = {"HIGH": "#FF5252", "MEDIUM": "#FFB74D", "LOW": "#66BB6A"}
        sev_order = ["HIGH", "MEDIUM", "LOW"]
        sev_counts = [
            sum(1 for t in threats if str(t.get("risk_level") or "").upper() == sev) for sev in sev_order
        ]
        sev_values = sev_counts
        sev_labels = sev_order

        # Severity donut/pie
        total_sev = sum(sev_values)
        if total_sev > 0:
            self.threat_ax_sev.set_title("Threat Severity Distribution", color=TEXT, fontsize=10)
            self.threat_ax_sev.axis("equal")
            wedges = self.threat_ax_sev.pie(
                sev_values,
                labels=sev_labels,
                colors=[color_map[s] for s in sev_order],
                startangle=90,
                autopct="%1.0f%%",
                textprops={"color": TEXT, "fontsize": 8},
                wedgeprops={"linewidth": 1, "edgecolor": "#2A2A2A", "width": 0.45},
            )
            # Avoid unused vars; matplotlib returns (wedges, texts, autotexts)
            _ = wedges
        else:
            self.threat_ax_sev.text(0.5, 0.5, "No threat data", ha="center", va="center", color=TEXT, transform=self.threat_ax_sev.transAxes)

        # Category breakdown (COUNT of threats per category)
        # Backend category "OS / Patching" is shown as "Patching" on the bar chart.
        cat_backend_map = {
            "Firewall": "Firewall",
            "Open Ports": "Open Ports",
            "Services": "Services",
            "Patching": "OS / Patching",
            "Antivirus": "Antivirus",
            "Third-party Software": "Third-party Software",
        }
        cat_order = ["Firewall", "Open Ports", "Services", "Patching", "Antivirus", "Third-party Software"]
        cat_counts = [
            sum(1 for t in threats if str(t.get("category") or "").strip() == cat_backend_map[label])
            for label in cat_order
        ]

        cat_color_map = {
            "Firewall": "#E53935",
            "Open Ports": "#FB8C00",
            "Services": "#8E24AA",
            "Patching": "#FDD835",
            "Antivirus": "#1E88E5",
            "Third-party Software": "#26A69A",
        }

        self.threat_ax_cat.set_title("Threat Distribution by Category", color=TEXT, fontsize=10)
        bars = self.threat_ax_cat.bar(
            cat_order,
            cat_counts,
            color=[cat_color_map[c] for c in cat_order],
            alpha=0.95,
        )
        self.threat_ax_cat.set_ylabel("Threat Count", color=TEXT, fontsize=8)
        self.threat_ax_cat.tick_params(axis="x", colors=TEXT, labelsize=8)
        self.threat_ax_cat.tick_params(axis="y", colors=TEXT, labelsize=8)

        for bar, val in zip(bars, cat_counts, strict=True):
            if val > 0:
                self.threat_ax_cat.text(
                    bar.get_x() + bar.get_width() / 2,
                    bar.get_height() + 0.05,
                    str(val),
                    ha="center",
                    va="bottom",
                    color=TEXT,
                    fontsize=8,
                )

        self.threat_figure.tight_layout(pad=1.0)
        self.threat_canvas.draw_idle()

    def _render_threat_analysis(self, scan: dict[str, Any]) -> None:
        if self.threat_inner_frame is None:
            return

        # Clear previous threats
        for child in self.threat_inner_frame.winfo_children():
            try:
                child.destroy()
            except Exception:
                pass

        try:
            threats = threat_analysis.generate_threats(scan, max_threats=10)
        except Exception as e:
            tk.Label(
                self.threat_inner_frame,
                text=f"Failed to generate Threat Analysis: {e}",
                bg=CARD,
                fg="#FF5252",
                wraplength=800,
                justify=tk.LEFT,
            ).pack(anchor="w", pady=6)
            return

        # Charts must reflect the exact same threats we render
        self._update_threat_charts(threats)

        # Group threats by category (keep a consistent, management-friendly order)
        order = ["Firewall", "Antivirus", "Open Ports", "OS / Patching", "Services", "Third-party Software"]
        grouped: dict[str, list[dict[str, Any]]] = {}
        for t in threats:
            cat = str(t.get("category") or "Other")
            grouped.setdefault(cat, []).append(t)

        color_map = {"HIGH": "#FF5252", "MEDIUM": "#FFB74D", "LOW": "#66BB6A"}

        any_added = False
        for cat in order:
            items = grouped.get(cat) or []
            if not items:
                continue
            any_added = True

            cat_lbl = tk.Label(
                self.threat_inner_frame,
                text=cat,
                bg=CARD,
                fg=TEXT,
                font=("Segoe UI", 13, "bold"),
                anchor="w",
                justify=tk.LEFT,
            )
            cat_lbl.pack(fill=tk.X, anchor="w", pady=(10, 8))

            for t in items:
                sev = str(t.get("risk_level") or "LOW").upper()
                fg = color_map.get(sev, TEXT)
                card = tk.Frame(self.threat_inner_frame, bg="#171717", highlightbackground="#2A2A2A", highlightthickness=1)
                card.pack(fill=tk.X, anchor="w", pady=6)

                title = str(t.get("threat_title") or "Threat")
                nist_func = str(t.get("nist_function") or "")
                desc = str(t.get("description") or "")
                rec = str(t.get("recommendation") or "")

                tk.Label(
                    card,
                    text=f"Title: {title}",
                    bg="#171717",
                    fg=TEXT,
                    font=("Segoe UI", 11, "bold"),
                    anchor="w",
                    justify=tk.LEFT,
                    wraplength=900,
                ).pack(fill=tk.X, padx=10, pady=(10, 0))

                tk.Label(
                    card,
                    text=f"Risk Level: {sev}",
                    bg="#171717",
                    fg=fg,
                    font=("Consolas", 11, "bold"),
                    anchor="w",
                    justify=tk.LEFT,
                ).pack(fill=tk.X, padx=10, pady=(6, 0))
                if nist_func:
                    tk.Label(
                        card,
                        text=f"NIST Function: {nist_func}",
                        bg="#171717",
                        fg=MUTED,
                        anchor="w",
                        justify=tk.LEFT,
                    ).pack(fill=tk.X, padx=10, pady=(6, 0))

                tk.Label(
                    card,
                    text=f"Description: {desc}",
                    bg="#171717",
                    fg=TEXT,
                    anchor="w",
                    justify=tk.LEFT,
                    wraplength=900,
                ).pack(fill=tk.X, padx=10, pady=(10, 0))

                tk.Label(
                    card,
                    text=f"Recommendation: {rec}",
                    bg="#171717",
                    fg=TEXT,
                    anchor="w",
                    justify=tk.LEFT,
                    wraplength=900,
                ).pack(fill=tk.X, padx=10, pady=(8, 12))

        if not any_added:
            tk.Label(
                self.threat_inner_frame,
                text="No threats were generated from the available local findings.",
                bg=CARD,
                fg=MUTED,
                wraplength=900,
                justify=tk.LEFT,
            ).pack(anchor="w", pady=10)

    def _export_threat_report_clicked(self) -> None:
        """
        Builds a dedicated Word threat report for the selected system.
        """
        if not self.selected_scan:
            messagebox.showwarning("Export Threat Report", "Select a system first.")
            return

        try:
            hostname = str(self.selected_scan.get("hostname") or "system").strip() or "system"
            safe_name = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in hostname)
            out_path = build_threat_word_report(self.selected_scan, output_path=Path(f"threat_report_{safe_name}.docx"))
            messagebox.showinfo("Threat Report Generated", f"Threat Analysis Word report saved to: {out_path}")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to generate Threat Analysis report: {e}")

    def _export_threats_to_excel_clicked(self) -> None:
        """
        Exports threats to an Excel file using append mode.
        - If the file does not exist, it is created.
        - If it exists, new rows are appended (no overwrite).
        """
        try:
            from openpyxl import Workbook, load_workbook
        except Exception:
            messagebox.showerror(
                "Missing Dependency",
                "openpyxl is not installed. Run: python -m pip install -r requirements.txt",
            )
            return

        if not self.results:
            messagebox.showwarning("Export to Excel", "No scan results found to export.")
            return

        compliance = "NIST SP 800-30 aligned risk scoring & NIST CSF 2.0 mapping"
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        out_path = Path("security_report.xlsx")
        sheet_name = "Threats"

        headers = [
            "Timestamp",
            "System Name",
            "Risk Level",
            "Risk Score",
            "Category",
            "Threat",
            "Threat Risk",
            "NIST Function",
            "Description",
            "Recommendation",
            "Compliance",
        ]

        try:
            if out_path.exists():
                wb = load_workbook(out_path)
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                else:
                    ws = wb.create_sheet(sheet_name)
                    ws.append(headers)
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = sheet_name
                ws.append(headers)

            # If the sheet is empty (e.g., created with no headers), ensure headers exist.
            if ws.max_row == 0:
                ws.append(headers)

            for scan in self.results:
                scan_hostname = scan.get("hostname") or "Unknown"
                for t in threat_analysis.generate_threats(scan, max_threats=10):
                    row = [
                        timestamp,
                        scan_hostname,
                        str(scan.get("risk_level") or "UNKNOWN").upper(),
                        int(scan.get("risk_score") or 0),
                        str(t.get("category") or ""),
                        str(t.get("threat_title") or ""),
                        str(t.get("risk_level") or "LOW").upper(),
                        str(t.get("nist_function") or ""),
                        str(t.get("description") or ""),
                        str(t.get("recommendation") or ""),
                        compliance,
                    ]
                    ws.append(row)

            # Installed applications export (registry-based, from scan_results.json).
            apps_sheet_name = "Installed Applications"
            apps_headers = ["System Name", "Application Name", "Version"]
            risky_sheet_name = "Risky Applications"
            risky_headers = ["System Name", "Application Name", "Version", "Risk Level", "Reason"]
            try:
                if apps_sheet_name in wb.sheetnames:
                    apps_ws = wb[apps_sheet_name]
                else:
                    apps_ws = wb.create_sheet(apps_sheet_name)
                    apps_ws.append(apps_headers)

                if risky_sheet_name in wb.sheetnames:
                    risky_ws = wb[risky_sheet_name]
                else:
                    risky_ws = wb.create_sheet(risky_sheet_name)
                    risky_ws.append(risky_headers)
            except Exception:
                # If something goes wrong with sheet handling, fail silently for apps export.
                apps_ws = None
                risky_ws = None

            if apps_ws is not None:
                if apps_ws.max_row == 0:
                    apps_ws.append(apps_headers)

                for scan in self.results:
                    scan_hostname = scan.get("hostname") or "Unknown"
                    apps = scan.get("installed_apps") or []

                    # Backward compatibility: if installed_apps isn't present in older scan JSON,
                    # fetch from local registry for export.
                    if not apps:
                        try:
                            from . import installed_apps as installed_apps_module

                            apps = installed_apps_module.get_installed_apps()
                        except Exception:
                            apps = []

                    for app in apps:
                        app_name = str(app.get("name") or "").strip()
                        if not app_name:
                            continue
                        app_version = str(app.get("version") or "N/A").strip() or "N/A"
                        apps_ws.append([scan_hostname, app_name, app_version])

            if risky_ws is not None:
                if risky_ws.max_row == 0:
                    risky_ws.append(risky_headers)

                for scan in self.results:
                    scan_hostname = scan.get("hostname") or "Unknown"
                    risky_apps = scan.get("risky_apps")
                    if not isinstance(risky_apps, list):
                        risky_apps = risk.filter_risky_apps(
                            risk.analyze_software_risk(scan.get("installed_apps") or [])
                        )

                    for item in risky_apps:
                        risky_ws.append(
                            [
                                scan_hostname,
                                str(item.get("name") or "").strip(),
                                str(item.get("version") or "N/A").strip() or "N/A",
                                str(item.get("risk_level") or item.get("risk") or "LOW").upper(),
                                str(item.get("reason") or "").strip(),
                            ]
                        )

            wb.save(str(out_path))
            messagebox.showinfo("Export Complete", f"Threat data appended to: {out_path}")
        except Exception as e:
            messagebox.showerror("Export Failed", f"Failed to export to Excel: {e}")


def main() -> None:
    app = EndpointDashboard()
    app.mainloop()


if __name__ == "__main__":
    main()
